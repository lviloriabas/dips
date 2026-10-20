"""Las extensiones se presentan al auditor para revisar su autorización."""

from datetime import datetime

import openpyxl
import pytest

from dips_audit.__main__ import main
from dips_audit.cargador import cargar_export, lote_desde_icam
from dips_audit.comparacion import comparar, exportar
from dips_audit.reglas import EXTENSION, auditar_dip, estado_final
from dips_audit.reporte import exportar_excel
from tests.conftest import ENCABEZADO_ENTRADA, crear_export, hacer_dip
from tests.test_barcode import _registro
from tests.test_comparacion import ENCABEZADO, _fila
from tests.test_integracion import FILAS


@pytest.mark.parametrize("estado, completo", [
    ("ACTV", None),
    ("COMPLETE", datetime(2026, 8, 20)),  # Cerrado antes del vencimiento original.
    ("COMPLETE", datetime(2026, 8, 24)),
])
@pytest.mark.parametrize("ok_manual", [False, True])
def test_extension_se_conserva_en_resultado_y_excel(tmp_path, auditor, estado, completo, ok_manual):
    fila = list(FILAS[3])
    fila[6], fila[11] = estado, datetime(2026, 8, 25, 23, 59)
    entrada = crear_export(tmp_path / "entrada.xlsx",
                           [[*fila, datetime(2026, 8, 21, 23, 59), completo]],
                           encabezado=[*ENCABEZADO_ENTRADA, "Orig Fault Due Date", "Complete Dt"])
    lote = cargar_export(entrada)
    if ok_manual:
        lote.datos.at[0, "auditoria_manual"] = ["Cierre: sin discrepancia"]
    resultado = auditor.auditar_lote(lote)
    actual = resultado.filas.iloc[0]
    assert actual["RESULTADO"] == actual["ESTADO IA"] == EXTENSION
    assert actual["ESTADO APERTURA"] == EXTENSION
    assert actual["ESTADO CIERRE"] == (EXTENSION if completo else "")
    assert "Revisar la extensión" in actual["VERIFICACIÓN MANUAL"]
    assert "[EXTENSIÓN]" in actual["HALLAZGOS"]
    assert "exten" not in actual["COMENTARIO SUGERIDO"].lower()
    assert not (resultado.hallazgos_auditoria_actual()["Severidad"] == "POSIBLE DISCREPANCIA").any()
    wb = openpyxl.load_workbook(exportar_excel(resultado, tmp_path / "salida.xlsx"))
    resumen = {f[0]: f[1] for f in wb["Resumen"].values if len(f) > 1}
    assert resumen[EXTENSION] == 1 and resumen["POSIBLE DISCREPANCIA"] == 0
    col = next(c.column for c in wb["Dashboard"][5] if c.value == EXTENSION)
    assert wb["Dashboard"].cell(6, col).value == 1
    columnas = {c.value: c.column for c in wb["Auditoría"][1]}
    celda = wb["Auditoría"].cell(2, columnas["RESULTADO"])
    assert celda.value == EXTENSION and celda.fill.fgColor.rgb != "00000000"
    columna_resultado = next(c.column for c in wb["Hallazgos"][1] if c.value == "Resultado")
    assert EXTENSION in {f[columna_resultado - 1] for f in wb["Hallazgos"].values}


@pytest.mark.parametrize("parametro, unidad, extension", [
    ("CDY", "DAY", "30"), ("HOURS", "HOUR", "2.5"), ("CYCLES", "CYCLE", "5"),
])
def test_detecta_extension_del_ta_icam_sin_cambio_de_fecha(tmp_path, ctx, parametro, unidad, extension):
    limite = f"\t\t{parametro}\t{unidad}\tCUSTOM\tAircraft 1000\t30\tAircraft 1030\t{extension}\tAircraft 1035\t2\t21-AUG-2026 23:59 UTC"
    lote = lote_desde_icam(_registro(limites=[limite]), tmp_path / "icam.json")
    intervalos = lote.datos.iloc[0]["icam_intervalos"]
    assert intervalos[0]["extension"] == extension
    dip = hacer_dip(icam_intervalos=intervalos)
    r = auditar_dip(dip, None, ctx)
    assert estado_final(r.hallazgos) == EXTENSION
    assert any(extension in h.mensaje for h in r.hallazgos if h.codigo == "dip_extendido")


@pytest.mark.parametrize("extension", [None, "", "0", "0.0", "-1"])
def test_sin_extension_positiva_no_cambia_el_resultado(ctx, extension):
    dip = hacer_dip(orig_due_date=datetime(2026, 8, 21, 23, 59),
                    icam_intervalos=[{"extension": extension}])
    r = auditar_dip(dip, None, ctx)
    assert not any(h.codigo == "dip_extendido" for h in r.hallazgos)
    assert estado_final(r.hallazgos) == "OK"


@pytest.mark.parametrize("estado", ["ACTV", "COMPLETE"])
def test_extension_vencida_conserva_discrepancia(catalogo, ctx, estado):
    dip = hacer_dip(estado=estado, due_date=datetime(2026, 8, 19, 23, 59),
                    orig_due_date=datetime(2026, 8, 18, 23, 59),
                    complete_date=datetime(2026, 8, 20, 10), usage_remaining=-1)
    r = auditar_dip(dip, catalogo.por_id("OXIGENO_TRIPULACION"), ctx)
    assert estado_final(r.hallazgos) == "POSIBLE DISCREPANCIA"
    assert any(h.codigo == "dip_extendido" and h.severidad == EXTENSION for h in r.hallazgos)


def test_extension_no_oculta_otra_discrepancia(catalogo, ctx):
    dip = hacer_dip(due_date=datetime(2026, 8, 25, 23, 59),
                    orig_due_date=datetime(2026, 8, 21, 23, 59), intervalo=6)
    r = auditar_dip(dip, catalogo.por_id("OXIGENO_TRIPULACION"), ctx)
    assert estado_final(r.hallazgos) == "POSIBLE DISCREPANCIA"
    assert any(h.codigo == "intervalo_excedido" for h in r.hallazgos)


def test_extension_no_es_discrepancia_en_comparacion_y_consola(tmp_path, auditor, capsys):
    fila = _fila("E1", "COMPLETE", datetime(2026, 8, 25, 23, 59),
                  datetime(2026, 8, 24), None, datetime(2026, 8, 24, 15))
    entrada = crear_export(tmp_path / "entrada.xlsx", [[*fila, datetime(2026, 8, 21, 23, 59)]],
                           encabezado=[*ENCABEZADO, "Orig Fault Due Date"])
    c = comparar(auditor, entrada)
    assert c.detalle.iloc[0]["Programa: resultado"] == EXTENSION
    assert c.detalle.iloc[0]["¿Coinciden?"].startswith(EXTENSION)
    assert "extensión" in c.detalle.iloc[0]["Cierre: motivos"].lower()
    assert exportar(c, tmp_path / "salida").exists()
    assert main(["auditar", str(entrada), "--sin-excel", "--fecha", "2026-08-24"]) == 0
    texto = " ".join(capsys.readouterr().out.split())
    assert "EXTENSIÓN 1 100%" in texto
    assert "POSIBLE DISCREPANCIA 0 0%" in texto


def test_extension_de_item_con_vencimiento_se_revisa_con_fecha_original(catalogo, ctx):
    entrada = next(e for e in catalogo.entradas if e.hasta_vencimiento and not e.inventario)
    dip = hacer_dip(nombre="ITEM DUE DATE 21-AUG-2026", due_date=datetime(2026, 8, 25, 23, 59),
                    orig_due_date=datetime(2026, 8, 21, 23, 59))
    r = auditar_dip(dip, entrada, ctx)
    assert not any(h.codigo == "due_date_posterior_a_vencimiento_item" for h in r.hallazgos)
    assert any(h.codigo == "dip_extendido" for h in r.hallazgos)
