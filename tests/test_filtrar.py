from openpyxl import Workbook, load_workbook
import pytest

from dips_audit.filtrar import filtrar, pendiente, ruta_salida

ENCABEZADO = ["Fault Barcode", "Fault Status", "DIP Open Audit User", "DIP Closed Audit User"]


def test_pendiente():
    assert pendiente("ACTV", " ", " ")
    assert pendiente("ACTV (Active)", None, "X")
    assert not pendiente("ACTV", "KCAMPOS", " ")
    assert pendiente("COMPLETE", "KCAMPOS", " ")
    assert not pendiente("COMPLETE", " ", "KCAMPOS")
    assert not pendiente("CANCEL", " ", " ")


@pytest.mark.parametrize("estado", [
    "IN WORK", " in  work (In Work) ", "PAUSE", " pause (Paused) ",
    "ACTV (Active)", "ACTIVE", "OPEN",
])
def test_estados_abiertos_pendientes_de_apertura(estado):
    assert pendiente(estado, None, "KCAMPOS")
    assert not pendiente(estado, "KCAMPOS", None)


def test_filtrar(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "DIPs Audit (55)"
    ws.append(ENCABEZADO)
    ws.append(["A1", "ACTV", " ", " "])
    ws.append(["A2", "ACTV", "KCAMPOS", " "])
    ws.append(["C1", "COMPLETE", "KCAMPOS", "KCAMPOS"])
    ws.append(["C2", "COMPLETE", "KCAMPOS", " "])
    ws.append(["P1", "PAUSE", " ", " "])
    ws.append(["P2", "PAUSE", "KCAMPOS", " "])
    ws.append(["W1", "IN WORK (In Work)", " ", "KCAMPOS"])
    ws.append(["W2", "IN WORK", "KCAMPOS", " "])
    ws.append(["X1", "CANCEL", " ", " "])
    ws.auto_filter.ref = "A1:D10"
    otra = wb.create_sheet("Notas")
    otra.append(["sin columnas"])
    entrada = tmp_path / "DIPs Audit (25SEP2026).xlsx"
    wb.save(entrada)

    salida, resumen = filtrar(entrada)

    assert salida == ruta_salida(entrada) == tmp_path / "filtered" / "DIPs Audit (25SEP2026) FILTERED.xlsx"
    assert not (tmp_path / salida.name).exists()
    original = load_workbook(entrada)
    assert original["DIPs Audit (55)"].max_row == 10
    original.close()
    assert resumen == {"DIPs Audit (55)": (4, 9)}
    ws = load_workbook(salida)["DIPs Audit (55)"]
    assert [r[0] for r in ws.iter_rows(min_row=2, values_only=True)] == ["A1", "C2", "P1", "W1"]
    assert ws.auto_filter.ref == "A1:D5"


def test_filtrar_reutiliza_carpeta_y_actualiza_copia(tmp_path):
    entrada = tmp_path / "reporte.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(ENCABEZADO)
    ws.append(["A1", "ACTV", None, None])
    wb.save(entrada)
    salida, _ = filtrar(entrada)
    otro = salida.parent / "otro.txt"
    otro.write_text("conservar", encoding="utf-8")
    ws.append(["C1", "COMPLETE", None, None])
    wb.save(entrada)

    repetida, resumen = filtrar(entrada)

    assert repetida == salida
    assert resumen == {"Sheet": (2, 2)}
    assert otro.read_text(encoding="utf-8") == "conservar"
    copia = load_workbook(salida)
    assert copia.active.max_row == 3
    copia.close()
