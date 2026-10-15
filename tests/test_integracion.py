"""Pruebas de punta a punta: Excel de entrada -> auditoría -> Excel de resultado."""

from datetime import date, datetime

import openpyxl
import pytest

from dips_audit.__main__ import main
from dips_audit.auditor import Auditor
from dips_audit.reporte import _flota_dashboard, exportar_excel
from tests.conftest import ENCABEZADO_ENTRADA, crear_export

AVION = "Boeing 737-8V3 - HP-1837CMP"
FILAS = [
    # OK
    ["T1", 28, "Others", "DIP - FUEL MICROBIAL GROWTH", 4, "DAY", "ACTV", "HP-1837CMP", AVION,
     "DIP", "28-DIP", datetime(2026, 8, 22, 23, 59), datetime(2026, 8, 18, 9, 0), None, 2],
    # Intervalo excedido (FDR es 30 días)
    ["T2", 31, "Others", "FDR - REPORT", 45, "DAY", "ACTV", "HP-1837CMP", AVION, "DIP", "31-DIP",
     datetime(2026, 10, 2, 23, 59), datetime(2026, 8, 18, 9, 0), None, 43],
    # Sin referencia en la tabla
    ["T3", 25, "Others", "DEFIBRILLATOR DUE DATE", 38, "DAY", "ACTV", "HP-1837CMP", AVION, "DIP",
     "25-DIP", datetime(2026, 9, 25, 23, 59), datetime(2026, 8, 18, 9, 0), None, 36],
    # El segundo duplica al primero: posible discrepancia
    ["T4", 35, "Others", "CREW OXYGEN 850 PSI", 3, "DAY", "ACTV", "HP-1837CMP", AVION, "DIP",
     "35-DIP", datetime(2026, 8, 21, 23, 59), datetime(2026, 8, 18, 9, 0), None, 2],
    ["T5", 35, "Others", "DIP-CREW OXYGEN BOTTLE", 3, "DAY", "ACTV", "HP-1837CMP", AVION, "DIP",
     "35-DIP", datetime(2026, 8, 21, 23, 59), datetime(2026, 8, 18, 11, 0), None, 2],
    # Fuera de alcance (Powerplant)
    ["T6", 71, "Powerplant", "ENGINE 1 VCD SEAL", 40, "DAY", "ACTV", "HP-1837CMP", AVION, "DIP",
     "71-DIP", datetime(2026, 9, 27, 23, 59), datetime(2026, 8, 18, 9, 0), None, 38],
]


@pytest.fixture
def export(tmp_path):
    # Como el export real, con un correlativo "Index" en la primera columna.
    return crear_export(tmp_path / "export_icam.xlsx", [[n, *f] for n, f in enumerate(FILAS, 1)],
                        encabezado=["Index", *ENCABEZADO_ENTRADA])


def test_auditoria_de_punta_a_punta(auditor_others, export, tmp_path):
    resultado = auditor_others.auditar_archivo(export)
    # Tres resultados: OK, SIN FUENTE (no está en la tabla) o POSIBLE DISCREPANCIA.
    estados = dict(zip(resultado.filas["Fault Barcode"], resultado.filas["RESULTADO"]))
    assert estados == {
        "T1": "OK",
        "T2": "POSIBLE DISCREPANCIA",
        "T3": "SIN FUENTE EN DIPS REFERENCES",
        "T4": "OK",
        "T5": "POSIBLE DISCREPANCIA",
        "T6": "FUERA DE ALCANCE",
    }
    h = resultado.hallazgos
    assert set(h[(h["Fault Barcode"] == "T5") & (h["Código"] == "duplicado")]["Severidad"]) == {
        "POSIBLE DISCREPANCIA"}
    comentarios = resultado.filas.set_index("Fault Barcode")["COMENTARIO SUGERIDO"]
    assert "Posible duplicado de T4." in comentarios["T5"]
    assert comentarios["T1"] == "T/A aperturado correctamente"

    salida = exportar_excel(resultado, tmp_path / "salida.xlsx")
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames == ["Dashboard", "Resumen", "Auditoría", "Hallazgos"]
    dashboard = wb["Dashboard"]
    assert len(dashboard._charts) >= 3            # resultado, DIPs por tipo, referencias
    assert dashboard["B6"].value == 5             # DIPs auditados (sin el Powerplant)
    ws = wb["Auditoría"]
    encabezado = [c.value for c in ws[1]]
    # Columnas originales primero (menos Index), con su nombre y orden; las del programa al final.
    originales = [c.value for c in openpyxl.load_workbook(export)["Export"][1]][1:]
    assert encabezado[:len(originales)] == originales
    agregadas = encabezado[len(originales):]
    assert agregadas[:2] == ["RESULTADO", "COMENTARIO SUGERIDO"] and agregadas[-1] == "Fila Excel"
    assert ws.cell(1, len(originales)).fill.fgColor.rgb != ws.cell(1, len(originales) + 1).fill.fgColor.rgb
    assert "TIPO DE DISCREPANCIA" not in encabezado and "ESTADO IA" not in encabezado
    assert [c.value for c in wb["Hallazgos"][1]].count("Tipo de discrepancia") == 0
    assert ws.max_row == len(FILAS) + 1
    # En las dos hojas de DIPs el barcode es la primera columna y la única fija.
    for hoja in (ws, wb["Hallazgos"]):
        assert hoja["A1"].value == "Fault Barcode" and hoja.freeze_panes == "B2"
    assert [c.value for c in wb["Hallazgos"][1]][-1] == "Fila Excel"
    # El archivo original no se modifica.
    assert openpyxl.load_workbook(export).sheetnames == ["Export"]


def test_casos_aceptados_se_muestran_como_ok(auditor_others, export, tmp_path):
    resultado = auditor_others.auditar_archivo(export)
    fila = resultado.filas.set_index("Fault Barcode").loc["T4"]
    assert fila["RESULTADO"] == "OK"
    assert "[OK]" in fila["HALLAZGOS"]
    assert "condición" not in fila["COMENTARIO SUGERIDO"].lower()
    reporte = exportar_excel(resultado, tmp_path / "salida.xlsx")
    libro = openpyxl.load_workbook(reporte)
    hoja = libro["Hallazgos"]
    columnas = {c.value: c.column for c in hoja[1]}
    aceptados = [fila for fila in hoja.iter_rows(min_row=2)
                 if fila[columnas["Código"] - 1].value == "variante_condicional"]
    assert aceptados
    for fila in aceptados:
        celda = fila[columnas["Resultado"] - 1]
        assert celda.value == "OK"
        assert celda.fill.fgColor.rgb == "00C9E7C9"


def test_dip_repetido_es_discrepancia_hasta_que_uno_se_cierra(auditor_others, tmp_path):
    """Dos DIPs del mismo ítem abiertos a la vez: posible discrepancia. Si uno de los dos ya
    se cerró a la hora de la auditoría, el repetido está corregido y queda OK."""
    encabezado = [*ENCABEZADO_ENTRADA, "Complete Dt"]
    primero, segundo = list(FILAS[3]), list(FILAS[4])

    def duplicado(cierre_del_primero):
        primero[6] = "COMPLETE" if cierre_del_primero else "ACTV"
        ruta = crear_export(tmp_path / f"export_{primero[6]}.xlsx",
                            [[*primero, cierre_del_primero], [*segundo, None]], encabezado=encabezado)
        resultado = auditor_others.auditar_archivo(ruta)
        h = resultado.hallazgos
        return (resultado.filas.set_index("Fault Barcode").loc["T5", "RESULTADO"],
                h[(h["Fault Barcode"] == "T5") & (h["Código"] == "duplicado")].iloc[0])

    estado, hallazgo = duplicado(None)
    assert (estado, hallazgo["Severidad"]) == ("POSIBLE DISCREPANCIA", "POSIBLE DISCREPANCIA")
    estado, hallazgo = duplicado(datetime(2026, 8, 19, 10, 0))
    assert (estado, hallazgo["Severidad"]) == ("OK", "OK")
    assert "T4 ya se cerró" in hallazgo["Detalle"]


def test_sn_repetido_y_duplicado_declarado_abiertos_son_discrepancia(auditor_others, tmp_path):
    """El mismo S/N en dos DIPs activos del avión y el DIP que el técnico declaró duplicado y
    sigue abierto también son DIPs repetidos."""
    encabezado = [*ENCABEZADO_ENTRADA, "Fault Description"]
    fak = ["T7", 25, "Others", "FAK 1 S/N 12345 DUE DATE", 30, "DAY", "ACTV", "HP-1837CMP", AVION,
           "DIP", "25-DIP", datetime(2026, 9, 17, 23, 59), datetime(2026, 8, 18, 9, 0), None, 28]
    filas = [[*fak, ""],
             ["T8", *fak[1:3], "FAK 2 S/N 12345 DUE DATE", *fak[4:], ""],
             [*FILAS[0], "DIP DUPLICADO DEL T40S009J3HP"]]
    resultado = auditor_others.auditar_archivo(
        crear_export(tmp_path / "export.xlsx", filas, encabezado=encabezado))
    h = resultado.hallazgos
    severidades = {(b, c): s for b, c, s in zip(h["Fault Barcode"], h["Código"], h["Severidad"])}
    assert severidades[("T7", "sn_repetido")] == severidades[("T8", "sn_repetido")] ==         "POSIBLE DISCREPANCIA"
    assert severidades[("T1", "duplicado_declarado")] == "POSIBLE DISCREPANCIA"
    fila = resultado.filas.set_index("Fault Barcode").loc["T1"]
    assert fila["RESULTADO"] == "POSIBLE DISCREPANCIA"
    assert fila["COMENTARIO SUGERIDO"] == "DIP duplicado (ver T40S009J3HP); sigue abierto."


@pytest.mark.parametrize("estado, remanente, esperado", [
    ("ACTV", 43, "POSIBLE DISCREPANCIA"),
    ("COMPLETE", 43, "CORREGIDO EN CIERRE"),
    ("COMPLETE", -1, "POSIBLE DISCREPANCIA"),
])
def test_resultado_corresponde_a_la_auditoria_actual(auditor, tmp_path, capsys, estado,
                                                    remanente, esperado):
    fila = list(FILAS[1])  # FDR abierto con un intervalo incorrecto.
    fila[6], fila[-1] = estado, remanente
    completo = (datetime(2026, 10, 3, 10, 0) if remanente < 0
                else datetime(2026, 8, 20, 10, 0)) if estado == "COMPLETE" else None
    entrada = crear_export(tmp_path / "entrada.xlsx", [[*fila, completo]],
                           encabezado=[*ENCABEZADO_ENTRADA, "Complete Dt"])
    resultado = auditor.auditar_archivo(entrada)
    actual = resultado.filas.iloc[0]
    assert actual["ESTADO APERTURA"] == "POSIBLE DISCREPANCIA"
    assert actual["RESULTADO"] == actual["ESTADO IA"] == esperado
    assert resultado.conteo_estados() == {esperado: 1}
    assert "intervalo_excedido" in set(resultado.hallazgos["Código"])

    if estado == "COMPLETE":
        assert actual["AUDITORÍA"] == "CIERRE"
        assert actual["ESTADO CIERRE"] == ("OK" if esperado == "CORREGIDO EN CIERRE" else esperado)
        assert "T/A con 45 DAY" not in actual["COMENTARIO SUGERIDO"]
        if esperado == "CORREGIDO EN CIERRE":
            assert actual["COMENTARIO SUGERIDO"] == ""
        else:
            assert "Cerrado" in actual["COMENTARIO SUGERIDO"]

    salida = exportar_excel(resultado, tmp_path / "salida.xlsx")
    wb = openpyxl.load_workbook(salida)
    columnas = {c.value: c.column for c in wb["Auditoría"][1]}
    assert wb["Auditoría"].cell(2, columnas["RESULTADO"]).value == esperado
    assert wb["Dashboard"]["E6"].value == (1 if esperado == "POSIBLE DISCREPANCIA" else 0)
    assert wb["Dashboard"]["N6"].value == 0
    assert wb["Dashboard"]["Q6"].value == (1 if esperado == "CORREGIDO EN CIERRE" else 0)
    if esperado == "CORREGIDO EN CIERRE":
        assert wb["Dashboard"]["Q5"].value == esperado
        resumen = {f[0]: f[1] for f in wb["Resumen"].values if len(f) > 1}
        assert resumen[esperado] == 1 and resumen["OK"] == 0
        assert esperado in [c.value for c in wb["Dashboard"]["W"]]
        assert wb["Auditoría"].cell(2, columnas["RESULTADO"]).fill.fgColor.rgb != "00000000"
    assert main(["auditar", str(entrada), "--sin-excel", "--fecha", "2026-10-06"]) == 0
    consola = " ".join(capsys.readouterr().out.split())
    assert ("Posibles discrepancias por tipo" in consola) is (esperado == "POSIBLE DISCREPANCIA")
    if esperado == "CORREGIDO EN CIERRE":
        assert "POSIBLE DISCREPANCIA 0 0%" in consola
        assert "CORREGIDO EN CIERRE 1 100%" in consola
    elif estado == "COMPLETE":
        assert "(1 hallazgos)" in consola
    else:
        assert "(2 hallazgos)" in consola  # intervalo incorrecto y DIP vencido.


@pytest.mark.parametrize("remanente, esperado", [(36, "SIN FUENTE EN DIPS REFERENCES"),
                                               (-1, "POSIBLE DISCREPANCIA")])
def test_cierre_conserva_sin_fuente_y_discrepancias_del_cierre(auditor, tmp_path, remanente,
                                                             esperado):
    fila = list(FILAS[2])
    fila[6], fila[-1] = "COMPLETE", remanente
    entrada = crear_export(tmp_path / "entrada.xlsx", [[*fila, datetime(2026, 8, 20, 10, 0)]],
                           encabezado=[*ENCABEZADO_ENTRADA, "Complete Dt"])
    actual = auditor.auditar_archivo(entrada).filas.iloc[0]
    assert actual["RESULTADO"] == actual["ESTADO CIERRE"] == esperado


def test_cierre_ok_sin_discrepancias_de_apertura_permanece_ok(auditor, tmp_path):
    fila = list(FILAS[0])
    fila[6] = "COMPLETE"
    fila[1] = 31  # ATA distinto: OK en apertura.
    entrada = crear_export(tmp_path / "entrada.xlsx", [[*fila, datetime(2026, 8, 20)]],
                           encabezado=[*ENCABEZADO_ENTRADA, "Complete Dt"])
    resultado = auditor.auditar_archivo(entrada)
    actual = resultado.filas.iloc[0]
    assert "ata_distinto_a_referencia" in set(resultado.hallazgos["Código"])
    assert actual["ESTADO APERTURA"] == actual["ESTADO CIERRE"] == actual["RESULTADO"] == "OK"


@pytest.mark.parametrize("remanente, marcas, esperado", [
    (43, ["Apertura: DISCREPANCIA (X)", "Cierre: sin discrepancia"], "OK"),
    (-1, ["Cierre: sin discrepancia (DIP DUPLICATED)"], "OK"),
    (43, ["Apertura: sin discrepancia"], "CORREGIDO EN CIERRE"),
    (-1, ["Cierre: DISCREPANCIA (X)"], "POSIBLE DISCREPANCIA"),
])
def test_cierre_auditado_ok_en_icam_queda_ok_con_los_hallazgos(auditor, tmp_path, remanente,
                                                               marcas, esperado):
    """DIP COMPLETE con la marca de cierre sin discrepancia en Technical References: queda OK
    aunque el programa encuentre discrepancias; estados y hallazgos se conservan."""
    from dips_audit.cargador import cargar_export

    fila = list(FILAS[1])  # FDR con intervalo incorrecto: POSIBLE DISCREPANCIA en apertura.
    fila[6], fila[-1] = "COMPLETE", remanente
    completo = datetime(2026, 10, 3, 10, 0) if remanente < 0 else datetime(2026, 8, 20, 10, 0)
    lote = cargar_export(crear_export(tmp_path / "entrada.xlsx", [[*fila, completo]],
                                      encabezado=[*ENCABEZADO_ENTRADA, "Complete Dt"]))
    lote.datos.at[lote.datos.index[0], "auditoria_manual"] = marcas
    resultado = auditor.auditar_lote(lote)
    actual = resultado.filas.iloc[0]
    assert actual["RESULTADO"] == actual["ESTADO IA"] == esperado
    assert actual["ESTADO APERTURA"] == "POSIBLE DISCREPANCIA"
    assert actual["ESTADO CIERRE"] == ("OK" if remanente > 0 else "POSIBLE DISCREPANCIA")
    assert "intervalo_excedido" in set(resultado.hallazgos["Código"])
    assert actual["HALLAZGOS"]
    assert ("auditoría manual de cierre" in actual["VERIFICACIÓN MANUAL"]) is (esperado == "OK")


def test_evaluacion_cuenta_corregidos_en_cierre_sin_marcarlos_como_discrepancia(auditor, tmp_path):
    from dips_audit.evaluacion import evaluar

    fila = list(FILAS[1])
    fila[6] = "COMPLETE"
    encabezado = [*ENCABEZADO_ENTRADA, "Complete Dt", "STATUS - Apertura",
                  "COMMENTS - Apertura", "AUDIT DATE - Apertura"]
    entrada = crear_export(tmp_path / "entrada.xlsx",
                           [[*fila, datetime(2026, 8, 20), "OK", "", datetime(2026, 8, 18)]],
                           encabezado=encabezado)
    evaluacion = evaluar(auditor, entrada, hoja_entrada="Export", hoja_veredictos="Export")
    assert evaluacion.matriz.loc["OK", "CORREGIDO EN CIERRE"] == 1
    assert evaluacion.matriz.to_numpy().sum() == len(evaluacion.comparacion) == 1
    assert not evaluacion.comparacion["IA DISCREPANCIA"].any()


def test_dashboard_sin_dips_en_alcance(config, export, tmp_path):
    auditor = Auditor({**config, "tipos_auditados": ["Avionics"]}, fecha_auditoria=date(2026, 8, 20))
    salida = exportar_excel(auditor.auditar_archivo(export), tmp_path / "salida.xlsx")
    dashboard = openpyxl.load_workbook(salida)["Dashboard"]
    assert dashboard["B6"].value == 0 and dashboard["T6"].value == len(FILAS)


def test_cli_auditar(export, tmp_path, capsys):
    salida = tmp_path / "cli.xlsx"
    codigo = main(["auditar", str(export), "--salida", str(salida), "--fecha", "2026-08-20"])
    assert codigo == 0 and salida.exists()
    salida_consola = capsys.readouterr().out
    assert "POSIBLE DISCREPANCIA" in salida_consola and "CRÍTICA" not in salida_consola


def test_cli_auditar_con_flota_desconocida(tmp_path):
    desconocida = list(FILAS[0])
    desconocida[0], desconocida[7], desconocida[8] = "T0", "", "CFM56 ENGINE"
    entrada = crear_export(tmp_path / "entrada.xlsx", [desconocida, FILAS[1]])
    salida = tmp_path / "salida.xlsx"
    assert main(["auditar", str(entrada), "--salida", str(salida), "--fecha", "2026-08-20"]) == 0
    wb = openpyxl.load_workbook(salida)
    assert wb["Auditoría"].max_row == 3
    assert "Sin flota" in [c.value for c in wb["Dashboard"]["W"]]


def test_por_defecto_audita_todos_los_tipos(auditor, export):
    resultado = auditor.auditar_archivo(export)
    estados = dict(zip(resultado.filas["Fault Barcode"], resultado.filas["RESULTADO"]))
    assert "FUERA DE ALCANCE" not in estados.values()
    # Powerplant: fila VCD SEAL de la tabla, que va en el motor; este DIP está en el avión.
    fila = resultado.filas.set_index("Fault Barcode").loc["T6"]
    assert fila["REFERENCIA ASOCIADA"] == "VCD SEAL"
    assert estados["T6"] == "POSIBLE DISCREPANCIA" and "indica ENGINE" in fila["HALLAZGOS"]


def test_cli_tipos_dip(export, capsys):
    codigo = main(["auditar", str(export), "--sin-excel", "--fecha", "2026-08-20",
                   "--tipos-dip", "others", "Avionics"])
    assert codigo == 0
    salida = " ".join(capsys.readouterr().out.split())
    assert "Tipos de DIP auditados: others, Avionics" in salida
    assert "el Excel no tiene DIPs de tipo Avionics" in salida
    assert "FUERA DE ALCANCE 1 17%" in salida and "(Powerplant 1)" in salida


def test_cli_archivo_inexistente(tmp_path):
    assert main(["auditar", str(tmp_path / "no_existe.xlsx"), "--sin-excel"]) == 1


@pytest.mark.parametrize("flota, avion, esperado", [
    ("MAX", "HP-9810CMP", "MAX 8"), ("MAX", "HP-9924CMP", "MAX 9"),
    ("NG-700", "HP-1369CMP", "700"), ("NG-800", "HP-1837CMP", "800"),
    ("MAX", "", "Sin flota"), ("NG", "HP-1837CMP", "Sin flota"), ("", "", "Sin flota"),
    (None, "HP-1837CMP", "Sin flota"), (float("nan"), "HP-1837CMP", "Sin flota"),
    ("MAX", None, "Sin flota"), ("MAX", float("nan"), "Sin flota"),
])
def test_flotas_del_dashboard(flota, avion, esperado):
    assert _flota_dashboard(flota, avion) == esperado
