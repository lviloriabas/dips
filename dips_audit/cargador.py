"""Lectura del Excel de entrada y normalización de columnas y fechas."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import pandas as pd

from .estados import normalizar_estado
from .referencias import normalizar_unidad

# Columna estándar -> nombres aceptados en el Excel (sin distinguir mayúsculas).
COLUMNAS = {
    "barcode": ["Fault Barcode"],
    "ata": ["ATA"],
    "tipo": ["Type DIP", "Area"],
    "nombre": ["Fault Name"],
    "descripcion": ["Fault Description", "Description"],
    "accion_correctiva": ["Corrective Action", "corrective_action", "Acción Correctiva"],
    "intervalo": ["DIP Interval"],
    "tipo_intervalo": ["Interval Type"],
    "estado": ["Fault Status"],
    "avion": ["Aircraft"],
    "parent_assy": ["Task Parent Assy"],
    "severidad": ["Fault Severity"],
    "deferral_ref": ["Deferral Reference"],
    "due_date": ["Fault Due Date"],
    "orig_due_date": ["Orig Fault Due Date"],
    "found_date": ["Found On Date", "Fault Found on Date"],
    "defer_date": ["Defer Dt", "Deferral Date", "Fault Deferral Date"],
    "scheduled_date": ["Scheduled Date"],
    "work_package": ["Work Package Barcode"],
    "usage_remaining": ["Usage Remaining"],
    "complete_date": ["Complete Dt", "Completed Date", "Fault Completed Date", "Completion Date"],
    "referencia_tecnica": ["Technical References", "Technical Reference", "technical_references",
                           "Referencias Técnicas", "Referencia Técnica", "Referencia Tecnica"],
    # RIR activo del DIP (formato "DIPs Audit (23SEP2026)"): uno por fila.
    "rir_barcode": ["Rir Barcode"],
    "rir_nombre": ["Rir Name"],
    "rir_intervalo": ["Rir Driver Interval"],
    "rir_tipo": ["Rir Driver Type"],
    "rir_restante": ["Rir Driver Usage Rem"],
    # Attachments del DIP al exportar el Excel: "OK"/"PENDING" y cuántos hay.
    "adjuntos_estado": ["Attachment Status"],
    "adjuntos_cantidad": ["Attach Docs"],
}
OBLIGATORIAS = ["barcode", "nombre", "intervalo", "tipo_intervalo", "estado", "avion"]


class ErrorCarga(Exception):
    pass


@dataclass
class Lote:
    ruta: Path
    hoja: str
    fila_encabezado: int
    original: pd.DataFrame          # columnas tal como vienen en el Excel de entrada
    datos: pd.DataFrame             # columnas estándar ya interpretadas
    notas: list[str] = field(default_factory=list)
    prueba: bool = False            # datos de prueba (carpeta pruebas/): los reportes lo dicen


# --------------------------------------------------------------------------
# Conversión de valores
# --------------------------------------------------------------------------

_FORMATOS_FECHA = [
    "%b/%d/%y %H:%M", "%b/%d/%Y %H:%M", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %H:%M",
    "%m/%d/%Y", "%d/%b/%Y", "%d-%b-%Y", "%d-%b-%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
    "%d/%b/%Y %H:%M", "%d-%b-%Y %H:%M",
]


def parse_fecha(valor: object) -> datetime | None:
    # pd.NaT es instancia de datetime: se descarta antes de esa comprobación.
    if valor is None or valor is pd.NaT or (isinstance(valor, float) and pd.isna(valor)):
        return None
    if isinstance(valor, pd.Timestamp):
        return None if pd.isna(valor) else valor.to_pydatetime()
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, time())
    texto = str(valor).replace("\xa0", " ").strip()
    if not texto:
        return None
    for fmt in _FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, fmt)
        except ValueError:
            continue
    return None


def parse_numero(valor: object) -> float | None:
    if valor is None or isinstance(valor, (datetime, date, time)):
        return None
    if isinstance(valor, (int, float)):
        return None if pd.isna(valor) else float(valor)
    texto = str(valor).replace("\xa0", " ").replace(",", "").strip()
    try:
        return float(texto)
    except ValueError:
        return None


def parse_texto(valor: object) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    return str(valor).replace("\xa0", " ").strip()


def flota_de_modelo(texto: str) -> str | None:
    """'Boeing 737-9 MAX9' -> MAX; '737-8V3' -> NG-800; '737-7V3' -> NG-700."""
    t = texto.upper()
    if "MAX" in t:
        return "MAX"
    m = re.search(r"737-(\d)", t)
    if m:
        return {"7": "NG-700", "8": "NG-800", "9": "NG-900"}.get(m.group(1))
    return None


def flota_de_matricula(texto: str) -> str | None:
    """HP-9804CMP / HP-9901CMP -> MAX (matrículas 98xx y 99xx); HP-1837CMP -> NG."""
    m = re.search(r"\bHP-?(\d{2})\d{2}", texto.upper())
    if not m:
        return None
    return "MAX" if m.group(1) in ("98", "99") else "NG"


def flota(avion: str, parent_assy: str) -> str | None:
    """La matrícula decide MAX o NG; el modelo del parent assy precisa el NG (NG-700...)."""
    por_matricula = flota_de_matricula(avion) or flota_de_matricula(parent_assy)
    por_modelo = flota_de_modelo(parent_assy)
    if por_matricula == "NG" and por_modelo and por_modelo.startswith("NG"):
        return por_modelo
    return por_matricula or por_modelo


def tipo_parent(texto: str) -> str:
    t = texto.upper()
    if re.search(r"\bENGINE\b|CFM|LEAP", t):
        return "ENGINE"
    if re.search(r"\bAPU\b", t):
        return "APU"
    if "BOEING" in t or re.search(r"\b737\b|737-", t):
        return "AIRCRAFT"
    return "OTRO" if t else ""


# --------------------------------------------------------------------------
# Lectura
# --------------------------------------------------------------------------

def _buscar_encabezado(filas: list[tuple]) -> int | None:
    for i, fila in enumerate(filas[:20]):
        valores = {parse_texto(v).lower() for v in fila}
        if "fault barcode" in valores and "fault name" in valores:
            return i
    return None


def cargar_export(ruta: Path, hoja: str | None = None) -> Lote:
    ruta = Path(ruta)
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    hojas = [hoja] if hoja else wb.sheetnames
    if hoja and hoja not in wb.sheetnames:
        raise ErrorCarga(f"La hoja '{hoja}' no existe. Hojas: {', '.join(wb.sheetnames)}")
    for nombre in hojas:
        filas = list(wb[nombre].iter_rows(values_only=True))
        i = _buscar_encabezado(filas)
        if i is not None:
            break
    else:
        wb.close()
        raise ErrorCarga("No se encontró una hoja con las columnas 'Fault Barcode' y 'Fault Name'.")
    wb.close()

    encabezado = [parse_texto(v) or f"Columna {n + 1}" for n, v in enumerate(filas[i])]
    # Encabezados repetidos: se numeran para no perder columnas.
    vistos: dict[str, int] = {}
    for n, e in enumerate(encabezado):
        if e in vistos:
            vistos[e] += 1
            encabezado[n] = f"{e} ({vistos[e]})"
        else:
            vistos[e] = 1
    col_barcode = next(n for n, e in enumerate(encabezado) if e.lower() == "fault barcode")
    cuerpo, numeros_fila, vistos_barcode, repetidos = [], [], set(), 0
    for n, f in enumerate(filas[i + 1:], start=i + 2):  # n = número de fila en Excel
        f = (list(f) + [None] * len(encabezado))[:len(encabezado)]
        barcode = parse_texto(f[col_barcode])
        if not barcode:
            continue
        # El mismo DIP puede venir en dos filas (p. ej. auditado dos veces): se audita una vez.
        if barcode in vistos_barcode:
            repetidos += 1
            continue
        vistos_barcode.add(barcode)
        cuerpo.append(f)
        numeros_fila.append(n)
    original = pd.DataFrame(cuerpo, columns=encabezado, index=numeros_fila, dtype=object)

    datos, notas = normalizar_lote(original)
    if repetidos:
        notas.append(f"{repetidos} fila(s) con un Fault Barcode repetido se omitieron "
                     "(se audita la primera fila de cada DIP).")
    return Lote(ruta=ruta, hoja=nombre, fila_encabezado=i + 1, original=original,
                datos=datos, notas=notas)


def _mapear_columnas(columnas: list[str]) -> dict[str, str]:
    minusculas = {c.lower().strip(): c for c in columnas}
    mapa = {}
    for estandar, alias in COLUMNAS.items():
        for a in alias:
            if a.lower() in minusculas:
                mapa[estandar] = minusculas[a.lower()]
                break
    return mapa


def _rirs_del_excel(v) -> list:
    """RIR activo de la fila (columnas Rir del Excel). Sin Rir Barcode: el DIP no tiene RIR activo."""
    from .rir import RIR, Limite, unidad

    barcode = parse_texto(v("rir_barcode"))
    if not barcode:
        return []
    rir = RIR(parse_texto(v("rir_nombre")), [barcode], ["ACTV"], fuente="Excel")
    valor, tipo = parse_numero(v("rir_intervalo")), parse_texto(v("rir_tipo"))
    if valor is not None and tipo:
        rir.limites = [Limite(valor, unidad(tipo), restante=parse_numero(v("rir_restante")))]
    return [rir]


def _adjuntos_del_excel(v, mapa: dict) -> int | None:
    """Cantidad de attachments según el Excel (Attach Docs; PENDING = ninguno). None = sin dato."""
    if "adjuntos_estado" not in mapa and "adjuntos_cantidad" not in mapa:
        return None
    cantidad = parse_numero(v("adjuntos_cantidad"))
    if cantidad is not None:
        return int(cantidad)
    estado = parse_texto(v("adjuntos_estado")).upper()
    return 0 if estado == "PENDING" else None


def normalizar_lote(original: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    mapa = _mapear_columnas(list(original.columns))
    faltantes = [c for c in OBLIGATORIAS if c not in mapa]
    if faltantes:
        raise ErrorCarga("Faltan columnas obligatorias en el Excel: " + ", ".join(
            COLUMNAS[c][0] for c in faltantes))
    notas: list[str] = []
    registros = []
    for fila_excel, fila in original.iterrows():
        def v(col: str) -> object:
            return fila[mapa[col]] if col in mapa else None

        due = parse_fecha(v("due_date"))
        found = parse_fecha(v("found_date"))
        intervalo = parse_numero(v("intervalo"))
        unidad = normalizar_unidad(v("tipo_intervalo"))
        correccion = ""
        # Fechas con día y mes invertidos por la configuración regional de Excel
        # (p. ej. 9-jul guardado como 7-sep). Se corrige solo si el resultado es
        # coherente con el due date.
        if found and due and found.date() > due.date() and found.day <= 12:
            try:
                invertida = found.replace(month=found.day, day=found.month)
            except ValueError:
                invertida = None
            if invertida and invertida <= due:
                correccion = (f"Fecha de apertura corregida de {found:%d-%b-%Y} a "
                              f"{invertida:%d-%b-%Y} (día/mes invertidos en el Excel)")
                found = invertida
        parent = parse_texto(v("parent_assy"))
        ata = parse_numero(v("ata"))
        rirs_excel = _rirs_del_excel(v) if "rir_barcode" in mapa else None
        registros.append({
            "fila": fila_excel,
            "barcode": parse_texto(v("barcode")),
            "ata": int(ata) if ata is not None else None,
            "tipo": parse_texto(v("tipo")),
            "nombre": parse_texto(v("nombre")),
            "descripcion": parse_texto(v("descripcion")),
            "descripcion_disponible": "descripcion" in mapa,
            "intervalo": intervalo,
            "tipo_intervalo_original": parse_texto(v("tipo_intervalo")),
            "unidad": unidad,
            "estado": normalizar_estado(parse_texto(v("estado"))),
            "avion": parse_texto(v("avion")),
            "parent_assy": parent,
            "parent_tipo": tipo_parent(parent),
            "severidad": parse_texto(v("severidad")).upper(),
            "deferral_ref": parse_texto(v("deferral_ref")),
            "due_date": due,
            "orig_due_date": parse_fecha(v("orig_due_date")),
            "found_date": found,
            "defer_date": parse_fecha(v("defer_date")),
            "scheduled_date": parse_fecha(v("scheduled_date")),
            "work_package": parse_texto(v("work_package")),
            "usage_remaining": parse_numero(v("usage_remaining")),
            "complete_date": parse_fecha(v("complete_date")),
            # None = el Excel no trae el dato; "" = el DIP no tiene referencia técnica.
            "referencia_tecnica": (parse_texto(v("referencia_tecnica"))
                                   if "referencia_tecnica" in mapa else None),
            "correccion_fecha": correccion,
            # True si hay con qué revisar la referencia técnica (descripción o columna de
            # referencias en el Excel, o datos de ICAM).
            "datos_referencia": any(c in mapa for c in
                                    ("descripcion", "referencia_tecnica", "accion_correctiva")),
            "intervalo_de_icam": False,
            "notas_descripcion": "",
            "adjuntos": None,          # lista de nombres si se leyó ICAM; None = sin dato
            "adjuntos_excel": _adjuntos_del_excel(v, mapa),  # cantidad según el Excel
            "uso_encontrado": None,    # ICAM: {"CYCLES": n, "HOURS": n} al encontrar la falla
            "uso_completado": None,    # ICAM: lo mismo, al cerrar el DIP
            # RIR del DIP (columnas Rir del Excel o Follow-on Tasks de ICAM); None = sin dato
            "rirs": rirs_excel,
            "rirs_excel": rirs_excel,  # se conserva al completar con ICAM
            "historico_rir": False,    # se leyó Show Historical Tasks
            "icam_intervalos": None,   # T/A en Scheduling de ICAM (valor, parámetro, Start)
            "parent_icam": "",
            # Auditoría de cierre (datos de ICAM); None = no se leyó ICAM.
            "accion_correctiva": (parse_texto(v("accion_correctiva"))
                                   if "accion_correctiva" in mapa else None),
            "sn_removidos": None,
            "sn_instalados": None,
            "analogos_pn": None,
            "cambios_deadline": None,
            "marcas_auditoria": None,
            "marcas_auditdip": None,   # marcas AUDITDIP: ya se auditó manualmente
            "auditoria_manual": None,  # lo que anotó el auditor en cada marca
            # Inventario de generators (PDF adjunto): {"archivos", "generadores", "errores"};
            # None = no se leyeron los PDF en ICAM.
            "inventario": None,
        })
    # dtype=object conserva los None: si no, pandas los vuelve NaT/NaN, que son
    # verdaderos en un if y pasan los controles "is None" de las reglas.
    datos = pd.DataFrame(registros, dtype=object)
    if datos.empty:
        return datos, notas

    # Flota por avión: la matrícula dice si es MAX o NG, y el parent assy (cuando
    # es el avión) precisa el modelo NG; se propaga a los DIPs cuyo parent es un motor.
    def _flota(r) -> str | None:
        return flota(r["avion"], r["parent_assy"] if r["parent_tipo"] == "AIRCRAFT" else "")
    flotas: dict[str, str] = {}
    for _, r in datos.iterrows():
        f = _flota(r)
        if f and r["avion"] and (r["avion"] not in flotas or f.startswith("NG-")):
            flotas[r["avion"]] = f
    # También esta columna conserva None, sin convertirlo en NaN al inferir texto.
    datos["flota"] = pd.Series(
        [flotas.get(r["avion"]) or _flota(r) for _, r in datos.iterrows()],
        index=datos.index, dtype=object)
    corregidas = int((datos["correccion_fecha"] != "").sum())
    if corregidas:
        notas.append(f"{corregidas} fecha(s) de apertura con día/mes invertidos fueron corregidas.")
    if "complete_date" not in mapa:
        notas.append("El Excel no trae fecha de cierre: no se puede verificar si los DIPs "
                     "COMPLETE se cerraron dentro del plazo.")
    if "descripcion" not in mapa:
        notas.append("El Excel no trae la descripción larga: la búsqueda de S/N y referencias "
                     "técnicas se hace solo sobre el Fault Name.")
    return datos, notas


def _dato_icam(r: dict, ficha=None) -> dict:
    """Lo que usa la auditoría de un DIP leído en ICAM (registro del extractor)."""
    from .icam import leer_ficha
    from .rir import agrupar

    secciones = (r.get("icam_detail") or r).get("sections", {})
    ficha = ficha or leer_ficha(secciones, parse_texto(r["barcode"]), parse_texto(r.get("page_title")))
    # Página de cada RIR (la lee el extractor): intervalo, Start y Parent Assembly.
    detalles = {i: leer_ficha(d.get("sections", {}), i) for i, d in (r.get("rir") or {}).items()}
    return {
        "referencia_tecnica": "; ".join(ficha.referencias_tecnicas),
        "descripcion": ficha.descripcion,
        "descripcion_disponible": True,
        "notas_descripcion": ficha.notas_descripcion,
        "adjuntos": ficha.adjuntos,
        "intervalos": [{"value": i.valor, "parameter": i.parametro, "unit": i.unidad,
                        "start": i.inicio, "extension": i.extension,
                        "deadline": i.deadline, "extended_deadline": i.deadline_extendido}
                       for i in ficha.intervalos],
        "uso_encontrado": ficha.uso_encontrado,
        "uso_completado": ficha.uso_completado,
        "rirs": agrupar(ficha.tareas_siguientes, detalles),
        "historico_rir": bool((secciones.get("FollowOnHistorico") or {}).get("text")),
        "parent_icam": ficha.parent_assembly,
        "accion_correctiva": ficha.accion_correctiva,
        "defer_date": ficha.diferido,
        "sn_removidos": ficha.sn_removidos,
        "sn_instalados": ficha.sn_instalados,
        "analogos_pn": {pn: d.get("analogos", []) for pn, d in (r.get("partes") or {}).items()
                        if d.get("verificado")},
        "cambios_deadline": ficha.cambios_deadline,
        "marcas_auditoria": ficha.marcas_auditoria,
        "marcas_auditdip": ficha.marcas_auditdip,
        "auditoria_manual": ficha.auditoria_manual,
        "fin_real": ficha.fin_real,
        "inventario": _inventario(r),
    }


def _inventario(r: dict) -> dict | None:
    """Generators de los PDF adjuntos que leyó el extractor; None si no se leyeron."""
    from .inventario import leer_inventario

    pdfs = r.get("pdf")
    if pdfs is None:
        return None
    generadores = [g for datos in pdfs.values() for g in leer_inventario(datos.get("campos") or {})]
    return {"archivos": list(pdfs) + list(r.get("pdf_errores") or {}),
            "generadores": generadores, "errores": dict(r.get("pdf_errores") or {})}


def _registros_icam(ruta: Path) -> dict[str, dict]:
    """Datos de ICAM por barcode: referencia técnica, descripción, notas de la descripción,
    intervalos y adjuntos.

    Acepta el JSON que llena el extractor (`data/icam.json`, todos los DIPs juntos), otro
    JSON de extracción (`dips_dataset.json` o `icam_dips.json`), una carpeta con un JSON
    por DIP (formato anterior) o un CSV/Excel con las columnas
    `barcode`, `technical_references`, `description` o `corrective_action`. El texto de las
    pestañas se interpreta con `icam.leer_ficha`.
    """
    import json

    from .extractor_icam import cargar_registros

    if ruta.is_dir() or ruta.suffix.lower() == ".json":
        if ruta.is_dir():
            registros = [json.loads(x.read_text(encoding="utf-8")) for x in sorted(ruta.glob("*.json"))]
        else:
            # JSON con "records" (todos los DIPs), una lista o un solo DIP. Si quedan JSON
            # sueltos en la carpeta del mismo nombre (p. ej. data/icam/), se unen antes.
            registros = list(cargar_registros(ruta).values())
            if not ruta.exists():
                raise ErrorCarga(f"No se encontraron los datos de ICAM: {ruta}")
        salida = {}
        for r in registros:
            if r.get("icam_capture_status", "six_tabs") != "six_tabs":
                continue  # solo vino del Excel: no hay datos de ICAM
            salida[parse_texto(r["barcode"])] = _dato_icam(r)
        return salida
    if ruta.suffix.lower() == ".csv":
        extra = pd.read_csv(ruta, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    else:
        extra = pd.read_excel(ruta, dtype=str).fillna("")
    columnas = {c.lower().strip(): c for c in extra.columns}
    col_barcode = columnas.get("barcode") or columnas.get("fault barcode")
    col_ref = next((columnas[a.lower()] for a in COLUMNAS["referencia_tecnica"]
                    if a.lower() in columnas), None)
    col_desc = columnas.get("description") or columnas.get("descripcion")
    col_accion = next((columnas[a.lower()] for a in COLUMNAS["accion_correctiva"]
                       if a.lower() in columnas), None)
    col_diferido = next((columnas[a.lower()] for a in COLUMNAS["defer_date"]
                         if a.lower() in columnas), None)
    if not col_barcode or not any((col_ref, col_desc, col_accion)):
        raise ErrorCarga(f"El archivo de ICAM {ruta.name} debe tener 'barcode' y al menos "
                         "'technical_references', 'description' o 'corrective_action'.")
    return {parse_texto(r[col_barcode]): {"referencia_tecnica": parse_texto(r[col_ref]) if col_ref else "",
                                          "descripcion": parse_texto(r[col_desc]) if col_desc else "",
                                          "descripcion_disponible": col_desc is not None,
                                          "notas_descripcion": "", "adjuntos": None,
                                          "intervalos": [],
                                          **({"accion_correctiva": parse_texto(r[col_accion])}
                                             if col_accion else {}),
                                          **({"defer_date": parse_texto(r[col_diferido])}
                                             if col_diferido else {})}
            for _, r in extra.iterrows()}


def _sin_utc(texto: str) -> str:
    return texto.replace(" UTC", "").strip()


def lote_desde_icam(r: dict, ruta: Path) -> Lote:
    """Un DIP armado solo con lo leído en ICAM (`r`: su registro en `ruta`, data/icam.json),
    sin Excel de entrada. Los datos del Excel salen de las pestañas de ICAM; el tipo de DIP
    (Area) no está en ICAM y queda vacío, y la matrícula solo se conoce si el Parent Assembly
    es el avión (en un motor, la flota sale del modelo: LEAP es MAX y CFM56, NG)."""
    from .icam import leer_ficha

    ruta = Path(ruta)
    f = leer_ficha((r.get("icam_detail") or r).get("sections", {}), parse_texto(r["barcode"]),
                   parse_texto(r.get("page_title")))
    m = re.match(r"^Fault Details - (.*?)(?: \[[^\]]*\])?$", f.titulo)
    matricula = re.search(r"\bHP-?\d{4}[A-Z]*\b", f.parent_assembly.upper())
    ata = re.match(r"\d{2}", f.deferral_reference)
    paquete = re.search(r"\[([A-Z0-9]+)\]\s*$", f.paquete_trabajo)
    # Due date, como el Excel de entrada: el límite que vence primero. Uno por fecha vence
    # en su Deadline; uno por uso (Deadline "Aircraft n / ..."), en su Estimated Date.
    intervalo = f.intervalos[0] if f.intervalos else None

    def vence(i, extendido: bool) -> datetime | None:
        fecha = parse_fecha(_sin_utc(i.deadline_extendido if extendido else i.deadline))
        return fecha or parse_fecha(_sin_utc(i.fecha_estimada))
    due = min(filter(None, (vence(i, True) for i in f.intervalos)), default=None)
    due_original = min(filter(None, (vence(i, False) for i in f.intervalos)), default=None)
    estado = normalizar_estado(f.estado_tarea or f.fault_status)
    fila = {
        "Fault Barcode": f.barcode,
        "ATA": ata.group(0) if ata else None,
        "Fault Name": m.group(1) if m else f.titulo,
        "Fault Description": f.descripcion,
        "Fault Severity": f.severidad.split(" (")[0],
        "Fault Status": estado,
        "Found On Date": _sin_utc(f.found_on),
        "Defer Dt": _sin_utc(f.diferido),
        "Fault Due Date": due,
        "Orig Fault Due Date": due_original,
        # Cierre: cuando se certificó la falla (el Complete Dt del Excel).
        "Complete Dt": _sin_utc(f.certificado or f.fin_real) if estado == "COMPLETE" else None,
        "DIP Interval": intervalo.valor if intervalo else None,
        "Interval Type": intervalo.parametro if intervalo else None,
        "Usage Remaining": intervalo.restante if intervalo else None,
        "Aircraft": matricula.group(0) if matricula else "",
        "Task Parent Assy": f.parent_assembly,
        "Deferral Reference": f.deferral_reference,
        "Scheduled Date": _sin_utc(f.programado),
        "Work Package Barcode": paquete.group(1) if paquete else "",
    }
    original = pd.DataFrame([fila], index=[1], dtype=object)
    datos, _ = normalizar_lote(original)
    if not datos.at[0, "flota"] and tipo_parent(f.parent_assembly) == "ENGINE":
        motor = f.parent_assembly.upper()
        datos.at[0, "flota"] = "MAX" if "LEAP" in motor or "MAX" in motor else (
            "NG" if "CFM56" in motor else None)
    lote = Lote(ruta=ruta, hoja="ICAM", fila_encabezado=0, original=original, datos=datos)
    completar_con_icam(lote, ruta, {f.barcode: _dato_icam(r, f)})
    lote.notas.clear()
    if not matricula:
        lote.notas.append("ICAM no trae la matrícula (el Parent Assembly no es el avión): no se "
                          "revisa si los attachments son de otro avión.")
    return lote


def completar_con_icam(lote: Lote, ruta: Path, registros: dict[str, dict] | None = None) -> int:
    """Agrega al lote los datos extraídos de ICAM, por barcode.

    - Referencia técnica y descripción: para revisar la referencia del manual.
    - Intervalo: solo si el Excel de entrada no lo trae (queda anotado en el DIP).
    Un barcode que no está en el archivo queda sin el dato y no se revisa.
    `registros`: los datos ya interpretados (si no, se leen de `ruta`).
    Devuelve cuántos DIPs del lote se completaron.
    """
    ruta = Path(ruta)
    registros = _registros_icam(ruta) if registros is None else registros
    datos = lote.datos
    completados = intervalos = 0
    for i, barcode in datos["barcode"].items():
        r = registros.get(barcode)
        if r is None:
            continue
        completados += 1
        datos.at[i, "datos_referencia"] = True
        datos.at[i, "referencia_tecnica"] = r["referencia_tecnica"]
        if r["descripcion"]:
            datos.at[i, "descripcion"] = r["descripcion"]
        if r.get("descripcion_disponible"):
            datos.at[i, "descripcion_disponible"] = True
        datos.at[i, "notas_descripcion"] = r["notas_descripcion"]
        datos.at[i, "adjuntos"] = r["adjuntos"]
        if r.get("rirs") is not None:  # ICAM manda; si no trae RIR, quedan los del Excel
            datos.at[i, "rirs"] = r["rirs"]
        datos.at[i, "historico_rir"] = r.get("historico_rir", False)
        datos.at[i, "icam_intervalos"] = r["intervalos"]
        datos.at[i, "parent_icam"] = r.get("parent_icam", "")
        if r.get("defer_date"):
            fecha_icam = parse_fecha(r["defer_date"].replace(" UTC", ""))
            if fecha_icam is not None:
                datos.at[i, "defer_date"] = fecha_icam
        for clave in ("accion_correctiva", "sn_removidos", "sn_instalados", "analogos_pn", "cambios_deadline",
                      "marcas_auditoria", "marcas_auditdip", "auditoria_manual", "uso_encontrado",
                      "uso_completado", "inventario"):
            if clave in r:
                datos.at[i, clave] = r[clave]
        # Fecha de cierre: si el Excel no la trae, la de ICAM (Actual End Date, en UTC).
        if datos.at[i, "complete_date"] is None and r.get("fin_real"):
            datos.at[i, "complete_date"] = parse_fecha(r["fin_real"].replace(" UTC", ""))
        if datos.at[i, "intervalo"] is None and r["intervalos"]:
            primero = r["intervalos"][0]
            valor = parse_numero(primero.get("value") or primero.get("interval"))
            unidad = normalizar_unidad(primero.get("parameter")) or normalizar_unidad(
                primero.get("unit"))
            if valor is not None and unidad:
                datos.at[i, "intervalo"] = valor
                datos.at[i, "unidad"] = unidad
                datos.at[i, "tipo_intervalo_original"] = parse_texto(primero.get("parameter"))
                datos.at[i, "intervalo_de_icam"] = True
                intervalos += 1
    lote.notas.append(f"Datos de ICAM ({ruta.name}): {completados} de {len(datos)} DIPs"
                      + (f"; intervalo tomado de ICAM en {intervalos}." if intervalos else "."))
    return completados
