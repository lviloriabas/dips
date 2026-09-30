"""Exporta el resultado de la auditoría a un Excel nuevo (el original no se modifica)."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.drawing.text import Font as DrawingFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .auditor import ResultadoAuditoria
from .reglas import MOSTRAR_SEVERIDAD

# El resultado es OK, REVISAR, SIN FUENTE o POSIBLE DISCREPANCIA: sin niveles (crítica / no
# crítica) ni tipos de discrepancia, que todavía no coinciden con el proceso de los auditores.
COLORES_RESULTADO = {"POSIBLE DISCREPANCIA": "FCE3B0", "REVISAR": "FFF59D", "SIN FUENTE": "D6E4F0",
                     "OK": "C9E7C9", "FUERA DE ALCANCE": "E4E4E4"}
COLORES_SEVERIDAD = {"POSIBLE DISCREPANCIA": "FCE3B0", "REVISAR": "FFF59D", "SIN FUENTE": "D6E4F0",
                     "OBSERVACIÓN": "EFEFEF"}
# Mismos tonos que las celdas, pero saturados (colores del tema de Office): en los gráficos los
# pasteles no se distinguen.
COLORES_GRAFICO = {"POSIBLE DISCREPANCIA": "ED7D31", "REVISAR": "FFC000", "SIN FUENTE": "4472C4",
                   "OK": "70AD47", "FUERA DE ALCANCE": "A5A5A5"}
_ENCABEZADO = PatternFill("solid", fgColor="1F3A5F")
# Columnas agregadas por el programa: encabezado morado y celdas lila, para
# distinguirlas de las columnas originales del Excel de entrada.
_ENCABEZADO_IA = PatternFill("solid", fgColor="5B3A8C")
_CELDA_IA = PatternFill("solid", fgColor="EEE8F6")
COLUMNAS_IA = ["RESULTADO", "COMENTARIO SUGERIDO", "HALLAZGOS",
               "VERIFICACIÓN MANUAL", "REFERENCIA ASOCIADA", "ASOCIACIÓN",
               "TA SEGÚN REFERENCIA", "REFERENCIA TÉCNICA ESPERADA", "REFERENCIA TÉCNICA DEL DIP",
               "RIR EN ICAM", "ADJUNTOS EN ICAM", "AUDITORÍA MANUAL EN ICAM", "AUDITORÍA",
               "FECHA AUDITORÍA", "ESTADO APERTURA", "ESTADO CIERRE"]
_ANCHOS = {"Fault Name": 45, "COMENTARIO SUGERIDO": 60, "HALLAZGOS": 70,
           "VERIFICACIÓN MANUAL": 60, "TA SEGÚN REFERENCIA": 45,
           "REFERENCIA TÉCNICA ESPERADA": 40, "REFERENCIA TÉCNICA DEL DIP": 40, "Detalle": 90,
           "REFERENCIA ASOCIADA": 32, "AUDITORÍA MANUAL EN ICAM": 40, "RESULTADO": 24,
           "ESTADO APERTURA": 24, "ESTADO CIERRE": 24}
_ENVOLVER = {"HALLAZGOS", "VERIFICACIÓN MANUAL", "COMENTARIO SUGERIDO", "TA SEGÚN REFERENCIA",
             "Detalle", "Fault Name", "REFERENCIA TÉCNICA ESPERADA", "AUDITORÍA MANUAL EN ICAM"}


def ruta_salida(resultado: ResultadoAuditoria, carpeta: Path) -> Path:
    carpeta.mkdir(parents=True, exist_ok=True)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    prefijo = "PRUEBA_" if resultado.lote.prueba else ""
    return carpeta / f"{prefijo}{resultado.lote.ruta.stem}_AUDITORIA_{marca}.xlsx"


AVISO_PRUEBA = "PRUEBA: datos de prueba, no es una auditoría real."


def _valor(v: object) -> object:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v


def _escribir_tabla(ws, df: pd.DataFrame, columnas_ia: set[str] | None = None,
                    col_color: str | None = None, colores: dict | None = None) -> None:
    columnas_ia = columnas_ia or set()
    ws.append(list(df.columns))
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _ENCABEZADO_IA if celda.value in columnas_ia else _ENCABEZADO
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    for fila in df.itertuples(index=False):
        ws.append([_valor(v) for v in fila])
    for n, nombre in enumerate(df.columns, start=1):
        letra = get_column_letter(n)
        ws.column_dimensions[letra].width = _ANCHOS.get(nombre, 16)
        envolver = nombre in _ENVOLVER
        for celda in ws[letra][1:]:
            celda.alignment = Alignment(wrap_text=envolver, vertical="top")
            if nombre in columnas_ia:
                celda.fill = _CELDA_IA
            if isinstance(celda.value, datetime):
                celda.number_format = "dd-mmm-yyyy hh:mm"
    if col_color and colores and col_color in df.columns:
        letra = get_column_letter(list(df.columns).index(col_color) + 1)
        for celda in ws[letra][1:]:
            color = colores.get(celda.value)
            if color:
                celda.fill = PatternFill("solid", fgColor=color)
                celda.font = Font(bold=True)
    # Fijas la fila de encabezados y las columnas hasta el barcode (la primera en las hojas
    # de DIPs), para no perder de vista el DIP al desplazarse a la derecha.
    fijas = list(df.columns).index("Fault Barcode") + 1 if "Fault Barcode" in df.columns else 1
    ws.freeze_panes = ws.cell(2, fijas + 1).coordinate
    ws.auto_filter.ref = ws.dimensions


def _colorear(ws, df: pd.DataFrame, columna: str, colores: dict) -> None:
    letra = get_column_letter(list(df.columns).index(columna) + 1)
    for celda in ws[letra][1:]:
        if celda.value in colores:
            celda.fill = PatternFill("solid", fgColor=colores[celda.value])


def exportar_excel(resultado: ResultadoAuditoria, ruta: Path) -> Path:
    wb = Workbook()
    _hoja_dashboard(wb.active, resultado)
    _hoja_resumen(wb.create_sheet("Resumen"), resultado)

    # Hoja principal: las columnas originales, en su orden y con su nombre (menos Index, un
    # correlativo del export), y al final las columnas que agrega el programa (con su color).
    original = resultado.lote.original.drop(columns=["Index"], errors="ignore")
    ia = resultado.filas.set_index("Fila Excel")[COLUMNAS_IA]
    combinado = pd.concat([original, ia], axis=1, join="inner")
    combinado["Fila Excel"] = combinado.index
    hoja = wb.create_sheet("Auditoría")
    _escribir_tabla(hoja, combinado, set(COLUMNAS_IA) | {"Fila Excel"}, "RESULTADO",
                    COLORES_RESULTADO)
    for columna in ("ESTADO APERTURA", "ESTADO CIERRE"):
        _colorear(hoja, combinado, columna, COLORES_RESULTADO)

    _escribir_tabla(wb.create_sheet("Hallazgos"), hallazgos_para_mostrar(resultado.hallazgos), None,
                    "Resultado", COLORES_SEVERIDAD)
    wb.save(ruta)
    return ruta


def hallazgos_para_mostrar(h: pd.DataFrame) -> pd.DataFrame:
    """Hallazgos con su resultado (posible discrepancia u observación), sin tipo de discrepancia."""
    h = h.drop(columns=["Tipo de discrepancia"], errors="ignore")
    h = h.assign(Severidad=h["Severidad"].map(MOSTRAR_SEVERIDAD).fillna(h["Severidad"]))
    h = h.rename(columns={"Severidad": "Resultado"})
    # Fila Excel al final, como en la hoja Auditoría: el barcode queda en la primera columna.
    return h[[c for c in h.columns if c != "Fila Excel"] + ["Fila Excel"]]


_RESULTADOS = ["POSIBLE DISCREPANCIA", "REVISAR", "SIN FUENTE", "OK"]
_GRIS_TEXTO, _GRIS_SUAVE, _GRIS_LINEA, _GRIS_FONDO = "404040", "7F7F7F", "D9D9D9", "F2F2F2"
_AZUL = "1F3A5F"

# Cuadrícula del dashboard: seis indicadores de 2 columnas separados por una columna angosta.
# Los gráficos se anclan a esas celdas (inicio y fin), así el espaciado es el mismo en Excel
# desktop y en Excel web aunque cambie el ancho en píxeles.
_COL0 = 2                       # B
_COLS_INDICADOR, _N_INDICADORES = 2, 6
_ULTIMA_COL = _COL0 + _N_INDICADORES * (_COLS_INDICADOR + 1) - 2          # R
_MITAD = _COL0 + 3 * (_COLS_INDICADOR + 1) - 1                             # J (separador)
_COL_DATOS = _ULTIMA_COL + 2                                               # T
_FILA_GRAFICOS, _ALTO_GRAFICO = 9, 16                                      # filas de 20 px


def _fuente(tam: float, color: str = _GRIS_TEXTO, negrita: bool = False) -> CharacterProperties:
    return CharacterProperties(sz=int(tam * 100), b=negrita, solidFill=color,
                               latin=DrawingFont(typeface="Calibri"))


def _texto(tam: float, color: str = _GRIS_TEXTO, negrita: bool = False) -> RichText:
    cp = _fuente(tam, color, negrita)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _titulo(texto: str) -> Title:
    cp = _fuente(12, _GRIS_TEXTO, True)
    parrafo = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=texto)])
    return Title(tx=Text(rich=RichText(p=[parrafo])), overlay=False)


def _linea(color: str) -> GraphicalProperties:
    return GraphicalProperties(ln=LineProperties(solidFill=color, w=9525))


def _etiquetas(valor: bool = False, porcentaje: bool = False, color: str = "FFFFFF",
               posicion: str | None = None) -> DataLabelList:
    """Solo el valor o el porcentaje. Toman el formato de la celda: los ceros no se muestran."""
    e = DataLabelList()
    e.showVal, e.showPercent = valor, porcentaje
    e.showSerName = e.showCatName = e.showLegendKey = e.showLeaderLines = False
    e.txPr = _texto(9, color, True)
    if posicion:
        e.dLblPos = posicion
    return e


def _estilo(grafico, titulo: str) -> None:
    """Estilo de gráfico de Office: fondo blanco, borde gris fino y esquinas rectas."""
    grafico.title = _titulo(titulo)
    grafico.roundedCorners = False
    grafico.graphical_properties = GraphicalProperties(solidFill="FFFFFF", ln=LineProperties(
        solidFill=_GRIS_LINEA, w=9525))
    if grafico.legend is not None:
        grafico.legend.txPr = _texto(9, _GRIS_TEXTO)
        grafico.legend.overlay = False


def _barras(titulo: str, categorias: int, apiladas: bool = False) -> BarChart:
    """Barras horizontales; con pocas categorías, más separadas para que no queden gruesas."""
    g = BarChart()
    g.type = "bar"
    if apiladas:
        g.grouping, g.overlap = "stacked", 100
    g.gapWidth = 60 if categorias > 4 else 150
    _estilo(g, titulo)
    # Eje de categorías: etiquetas grises, sin marcas. Eje de valores oculto: los valores van
    # en las barras.
    g.x_axis.delete = False
    g.x_axis.majorTickMark = "none"
    g.x_axis.txPr = _texto(9, _GRIS_TEXTO)
    g.x_axis.spPr = _linea(_GRIS_LINEA)
    g.x_axis.majorGridlines = None
    g.x_axis.scaling.orientation = "maxMin"         # primera categoría arriba
    g.y_axis.delete = True
    g.y_axis.majorGridlines = None
    return g


def _anclar(ws, grafico, fila: int, col_ini: int, col_fin: int, filas: int) -> None:
    """Ancla el gráfico de la celda (fila, col_ini) hasta el fin de (fila + filas - 1, col_fin)."""
    grafico.anchor = TwoCellAnchor(_from=AnchorMarker(col=col_ini - 1, row=fila - 1),
                                   to=AnchorMarker(col=col_fin, row=fila + filas - 1))
    ws.add_chart(grafico)


def _tabla_datos(ws, fila: int, encabezado: list[str], filas: list[list]) -> tuple[int, int]:
    """Tabla de datos de un gráfico, a la derecha. Devuelve (fila del encabezado, última fila).
    Los ceros se ven en blanco, en la tabla y en las etiquetas del gráfico."""
    for j, e in enumerate(encabezado):
        c = ws.cell(fila, _COL_DATOS + j, e)
        c.font = Font(bold=True, color=_GRIS_TEXTO)
        c.border = Border(bottom=Side(style="thin", color=_AZUL))
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")
    for i, f in enumerate(filas, start=1):
        for j, v in enumerate(f):
            c = ws.cell(fila + i, _COL_DATOS + j, v)
            c.font = Font(color=_GRIS_TEXTO)
            c.border = Border(bottom=Side(style="thin", color=_GRIS_LINEA))
            if j:
                c.number_format = "0;-0;;@"
    return fila, fila + len(filas)


def _indicador(ws, i: int, nombre: str, n: int, pct: float | None, color: str) -> None:
    col = _COL0 + i * (_COLS_INDICADOR + 1)
    fondo = PatternFill("solid", fgColor=_GRIS_FONDO)
    for fila in (5, 6, 7):
        for c in range(col, col + _COLS_INDICADOR):
            celda = ws.cell(fila, c)
            celda.fill = fondo
            if fila == 5:
                celda.border = Border(top=Side(style="thick", color=color))
    ws.cell(5, col, nombre).font = Font(bold=True, size=9, color=_GRIS_SUAVE)
    ws.cell(6, col, n).font = Font(bold=True, size=24, color=_GRIS_TEXTO)
    if pct is not None:
        ws.cell(7, col, pct).number_format = "0.0%"
    ws.cell(7, col).font = Font(size=9, color=_GRIS_SUAVE)
    for fila in (5, 6, 7):
        ws.cell(fila, col).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + _COLS_INDICADOR - 1)


def _apiladas_por(ws, datos, columna: str, titulo: str, fila_tabla: int, fila: int,
                  col_ini: int, col_fin: int, encabezado: str, ordenar: bool = True,
                  orden: list[str] | None = None) -> int:
    """Barras apiladas por resultado para cada valor de `columna` (en el `orden` dado, o de mayor
    a menor si `ordenar`). Devuelve la próxima fila libre de la tabla de datos."""
    tabla = datos.groupby([columna, "RESULTADO"]).size().unstack(fill_value=0)
    if tabla.empty:
        return fila_tabla
    columnas = [r for r in _RESULTADOS + ["FUERA DE ALCANCE"] if r in tabla.columns]
    if orden:
        tabla = tabla.loc[[o for o in orden if o in tabla.index]
                          + [t for t in tabla.index if t not in orden]]
    elif ordenar:
        tabla = tabla.loc[tabla[columnas].sum(axis=1).sort_values(ascending=False).index]
    f1, f2 = _tabla_datos(ws, fila_tabla, [encabezado] + columnas,
                          [[str(t)] + [int(tabla.loc[t, c]) for c in columnas] for t in tabla.index])
    g = _barras(titulo, len(tabla), apiladas=True)
    g.add_data(Reference(ws, min_col=_COL_DATOS + 1, max_col=_COL_DATOS + len(columnas),
                         min_row=f1, max_row=f2), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=_COL_DATOS, min_row=f1 + 1, max_row=f2))
    for serie, c in zip(g.series, columnas):
        serie.graphicalProperties.solidFill = COLORES_GRAFICO[c]
        serie.graphicalProperties.line.solidFill = "FFFFFF"
        serie.dLbls = _etiquetas(valor=True, color=_GRIS_TEXTO if c == "REVISAR" else "FFFFFF")
    g.legend.position = "t"
    _anclar(ws, g, fila, col_ini, col_fin, _ALTO_GRAFICO)
    return f2 + 2


# Flotas del dashboard, en este orden. Los NG-900 salen como "900" y lo que no se pudo
# precisar (un MAX sin matrícula, un NG sin modelo) como "Sin flota".
_FLOTAS_DASHBOARD = ["MAX 8", "MAX 9", "700", "800"]


def _flota_dashboard(flota: str, avion: str) -> str:
    """MAX + HP-98xx -> MAX 8, MAX + HP-99xx -> MAX 9; NG-700 -> 700, NG-800 -> 800."""
    if flota == "MAX":
        m = re.search(r"\bHP-?(98|99)\d{2}", avion.upper())
        return f"MAX {m.group(1)[1]}" if m else "Sin flota"
    if flota.startswith("NG-"):
        return flota[3:]
    return "Sin flota"


def _top(ws, conteo, titulo: str, encabezado: str, fila_tabla: int, fila: int, col_ini: int,
         col_fin: int, filas: int) -> int:
    """Barras horizontales con las categorías más frecuentes. Devuelve la próxima fila libre."""
    if conteo.empty:
        return fila_tabla
    f1, f2 = _tabla_datos(ws, fila_tabla, [encabezado, "DIPs"],
                          [[str(k), int(v)] for k, v in conteo.items()])
    g = _barras(titulo, len(conteo))
    g.legend = None
    g.add_data(Reference(ws, min_col=_COL_DATOS + 1, min_row=f1, max_row=f2), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=_COL_DATOS, min_row=f1 + 1, max_row=f2))
    g.series[0].graphicalProperties.solidFill = COLORES_GRAFICO["POSIBLE DISCREPANCIA"]
    g.series[0].dLbls = _etiquetas(valor=True, color=_GRIS_TEXTO, posicion="outEnd")
    _anclar(ws, g, fila, col_ini, col_fin, filas)
    return f2 + 2


def _hoja_dashboard(ws, resultado: ResultadoAuditoria) -> None:
    """Primera hoja: indicadores y gráficos del resultado."""
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = _AZUL
    filas = resultado.filas
    tipos = dict(zip(resultado.lote.datos["fila"], resultado.lote.datos["tipo"]))
    datos = filas.assign(TIPO=[(tipos.get(f) or "Sin tipo") for f in filas["Fila Excel"]])
    en_alcance = datos[datos["RESULTADO"] != "FUERA DE ALCANCE"]
    total = len(en_alcance)

    # Anchos y altos de la cuadrícula
    ws.column_dimensions["A"].width = 2
    for c in range(_COL0, _ULTIMA_COL + 1):
        separador = (c - _COL0) % (_COLS_INDICADOR + 1) == _COLS_INDICADOR
        ws.column_dimensions[get_column_letter(c)].width = 2 if separador else 11
    ws.column_dimensions[get_column_letter(_ULTIMA_COL + 1)].width = 4
    ws.column_dimensions[get_column_letter(_COL_DATOS)].width = 36
    for c in range(_COL_DATOS + 1, _COL_DATOS + 6):
        ws.column_dimensions[get_column_letter(c)].width = 22
    for fila, alto in {1: 8, 2: 30, 3: 18, 4: 12, 5: 20, 6: 36, 7: 20, 8: 16}.items():
        ws.row_dimensions[fila].height = alto

    # Encabezado: título, archivo y fecha
    prueba = resultado.lote.prueba
    ws["B2"] = "Auditoría preliminar de DIPs"
    ws["B2"].font = Font(bold=True, size=18, color="C00000" if prueba else _AZUL)
    ws["B2"].alignment = Alignment(vertical="center")
    ws["B3"] = resultado.lote.ruta.name
    ws["B3"].font = Font(color=_GRIS_SUAVE)
    # La fecha es un número: no se desborda a la celda vecina, va en el ancho del último indicador.
    col_fecha = _ULTIMA_COL - _COLS_INDICADOR + 1
    fecha = ws.cell(3, col_fecha, datetime.now().replace(second=0, microsecond=0))
    fecha.number_format = "dd-mmm-yyyy hh:mm"
    fecha.font = Font(color=_GRIS_SUAVE)
    fecha.alignment = Alignment(horizontal="right")
    ws.merge_cells(start_row=3, start_column=col_fecha, end_row=3, end_column=_ULTIMA_COL)
    if prueba:
        aviso = ws.cell(2, _ULTIMA_COL, AVISO_PRUEBA)
        aviso.font = Font(bold=True, color="C00000")
        aviso.alignment = Alignment(horizontal="right", vertical="center")

    # Indicadores
    _indicador(ws, 0, "DIPs AUDITADOS", total, None, _AZUL)
    for i, r in enumerate(_RESULTADOS, start=1):
        n = int((en_alcance["RESULTADO"] == r).sum())
        _indicador(ws, i, r, n, n / total if total else 0, COLORES_GRAFICO[r])
    _indicador(ws, 5, "FUERA DE ALCANCE", int((datos["RESULTADO"] == "FUERA DE ALCANCE").sum()),
               None, COLORES_GRAFICO["FUERA DE ALCANCE"])

    izquierda, derecha = (_COL0, _MITAD - 1), (_MITAD + 1, _ULTIMA_COL)
    fila1 = _FILA_GRAFICOS
    fila2 = fila1 + _ALTO_GRAFICO + 1
    fila3 = fila2 + _ALTO_GRAFICO + 1

    # Resultado (anillo), sin los resultados en cero: su etiqueta de % no se puede ocultar.
    conteo = {r: int((en_alcance["RESULTADO"] == r).sum()) for r in _RESULTADOS}
    presentes = [r for r in _RESULTADOS if conteo[r]] or _RESULTADOS
    f1, f2 = _tabla_datos(ws, 5, ["Resultado", "DIPs"], [[r, conteo[r]] for r in presentes])
    anillo = DoughnutChart(holeSize=58)
    anillo.add_data(Reference(ws, min_col=_COL_DATOS + 1, min_row=f1, max_row=f2), titles_from_data=True)
    anillo.set_categories(Reference(ws, min_col=_COL_DATOS, min_row=f1 + 1, max_row=f2))
    for i, r in enumerate(presentes):
        punto = DataPoint(idx=i)
        punto.graphicalProperties.solidFill = COLORES_GRAFICO[r]
        punto.graphicalProperties.line.solidFill = "FFFFFF"
        anillo.series[0].dPt.append(punto)
    anillo.dataLabels = _etiquetas(porcentaje=True)
    _estilo(anillo, "Resultado")
    anillo.legend.position = "r"
    _anclar(ws, anillo, fila1, *izquierda, _ALTO_GRAFICO)
    siguiente = f2 + 2

    # Apertura y cierre, tipo de DIP y flota: barras apiladas por resultado
    siguiente = _apiladas_por(ws, en_alcance, "AUDITORÍA", "Apertura y cierre", siguiente,
                              fila1, *derecha, "Auditoría", ordenar=False)
    siguiente = _apiladas_por(ws, datos, "TIPO", "DIPs por tipo", siguiente, fila2, *izquierda,
                              "Tipo DIP")
    flotas = en_alcance.assign(Flota=[_flota_dashboard(f, a) for f, a in
                                      zip(en_alcance["Flota"], en_alcance["Aircraft"])])
    siguiente = _apiladas_por(ws, flotas, "Flota", "DIPs por flota", siguiente, fila2, *derecha,
                              "Flota", orden=_FLOTAS_DASHBOARD)

    # Referencias y aviones con más posibles discrepancias
    discrepancias = en_alcance[en_alcance["RESULTADO"] == "POSIBLE DISCREPANCIA"]
    refs = discrepancias["REFERENCIA ASOCIADA"].replace("", "Sin referencia").value_counts().head(10)
    aviones = discrepancias["Aircraft"].replace("", "Sin avión").value_counts().head(10)
    alto = max(_ALTO_GRAFICO, 2 * max(len(refs), len(aviones)) + 5)
    siguiente = _top(ws, refs, "Referencias con posible discrepancia", "Referencia", siguiente,
                     fila3, *izquierda, alto)
    _top(ws, aviones, "Aviones con posible discrepancia", "Aircraft", siguiente, fila3, *derecha,
         alto)

    # Impresión: el dashboard en una página, sin las tablas de datos
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(_ULTIMA_COL + 1)}{fila3 + alto}"




def _hoja_resumen(ws, resultado: ResultadoAuditoria) -> None:
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90
    ws.append(["Auditoría preliminar de DIPs"])
    ws["A1"].font = Font(bold=True, size=14)
    if resultado.lote.prueba:
        ws.append([AVISO_PRUEBA])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color="C00000")
    ws.append([])
    datos = [
        ("Archivo", str(resultado.lote.ruta)),
        ("Hoja", resultado.lote.hoja),
        ("Generado", datetime.now().strftime("%d-%b-%Y %H:%M")),
        ("DIPs en el archivo", len(resultado.filas)),
        ("Tiempo de procesamiento", f"{resultado.segundos:.1f} s"),
    ]
    for fila in datos:
        ws.append(list(fila))
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["Resultado", "DIPs"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    conteo = resultado.filas["RESULTADO"].value_counts()
    for valor, color in COLORES_RESULTADO.items():
        ws.append([valor, int(conteo.get(valor, 0))])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=color)
    if resultado.advertencias:
        ws.append([])
        ws.append(["Notas de la carga"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for nota in resultado.advertencias:
            ws.append(["", nota])
            ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True)
