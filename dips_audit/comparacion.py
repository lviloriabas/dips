"""Compara el programa con la auditoría manual del Excel "DIPs Audit (...)".

Cada DIP se audita dos veces: al abrirse (apertura, DIP ACTV) y al cerrarse (cierre, DIP
COMPLETE). El Excel trae el día y la hora de cada auditoría manual ("DIP Open - RT Audit
Completed" y "DIP Closed - RT Audit Completed"), quién la hizo y si se encontró discrepancia
("DIP Discrepancy Found" = Y, una marca por DIP, no por auditoría).

El programa audita cada DIP a la fecha y hora de cada auditoría manual: con la de apertura
cuenta los hallazgos de apertura, y con la de cierre los de cierre. El resultado del programa
para el DIP es el peor de las auditorías que hizo el auditor.

Limitación: ICAM muestra el estado actual del DIP, no el que vio el auditor en ese momento
(p. ej. un RIR ya cerrado o un attachment agregado después).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .auditor import FUERA_DE_ALCANCE, Auditor
from .cargador import ErrorCarga, Lote, cargar_export, completar_con_icam, parse_fecha, parse_texto
from .reglas import SEVERIDAD_SIN_FUENTE, SIN_FUENTE

COLUMNAS_MANUALES = {
    "apertura": "DIP Open - RT Audit Completed",
    "auditor_apertura": "DIP Open Audit User",
    "cierre": "DIP Closed - RT Audit Completed",
    "auditor_cierre": "DIP Closed Audit User",
    "discrepancia": "DIP Discrepancy Found",
}
# Estado del programa -> texto para el reporte, de mejor a peor. El programa da cuatro
# resultados: OK, sin fuente (el DIP no está en la tabla de referencias), revisar (la tabla
# da el TA o el RIR según el daño: ir al manual) o posible discrepancia (sin niveles).
POSIBLE = "Posible discrepancia"
REVISAR = "Revisar en el manual"
RESULTADOS = ["OK", SIN_FUENTE, REVISAR, POSIBLE]
# Gravedad interna del hallazgo -> cómo se muestra.
SEVERIDADES = {"CRITICA": POSIBLE, "NO CRITICA": POSIBLE, "REVISAR": POSIBLE,
               "CONSULTAR MANUAL": REVISAR, SEVERIDAD_SIN_FUENTE: SIN_FUENTE, "OBSERVACION": "Observación"}
# Código del hallazgo -> motivo en palabras simples.
MOTIVOS = {
    "sin_referencia": f"{SIN_FUENTE}: el DIP no está en la tabla de referencias",
    "sin_referencia_tecnica": "No tiene referencia técnica (ni en ICAM ni en la descripción)",
    "referencia_tecnica_incompleta": "Referencia técnica sin revisión o sin fecha",
    "referencia_tecnica_equivocada": "Referencia técnica de otro ítem de la tabla",
    "referencia_tecnica_distinta": "Referencia técnica que no está en la tabla",
    "referencia_tecnica_mismo_tema": "Referencia técnica del mismo tema que la tabla, otra tarea",
    "intervalo_excedido": "Intervalo mayor que el permitido por la tabla",
    "intervalo_menor": "Intervalo menor que el de la tabla",
    "intervalo_cero": "Intervalo en cero",
    "tipo_intervalo_incorrecto": "Tipo de intervalo distinto (días, horas, ciclos)",
    "variante_condicional": "Intervalo de una condición especial de la tabla",
    "parent_assy_incorrecto": "Parent Assembly distinto del de la tabla",
    "due_date_posterior_a_intervalo": "Due date posterior a apertura + intervalo",
    "due_date_anterior_a_intervalo": "Due date anterior a apertura + intervalo",
    "due_date_posterior_a_vencimiento_item": "Due date posterior al vencimiento del ítem",
    "due_date_anterior_a_vencimiento_item": "Due date anterior al vencimiento del ítem",
    "sn_fuera_de_inventario": "S/N que no está en el PDF de inventario de generators",
    "sin_inventario_generadores": "Sin el PDF de inventario de generators",
    "dip_vencido": "DIP vencido a la hora de la auditoría",
    "usage_remaining_mayor_a_intervalo": "Usage Remaining mayor que el intervalo",
    "deferral_reference_ata": "Deferral Reference no coincide con el ATA",
    "ata_distinto_a_referencia": "ATA distinto del de la tabla",
    "tipo_dip_distinto_a_referencia": "Parece un DIP de otro tipo (Powerplant/Structures)",
    "sin_sn_en_descripcion": "Sin S/N del componente en la descripción",
    "sn_repetido": "Mismo S/N en otro DIP activo del avión",
    "duplicado": "Posible DIP duplicado",
    "duplicado_declarado": "DIP duplicado o creado por error",
    "fault_severity": "Fault Severity distinta de DIP",
    "intervalo_de_icam": "Intervalo tomado de ICAM (el Excel no lo trae)",
    "rir_faltante": "Falta el RIR que pide la tabla",
    "rir_intervalo_incorrecto": "RIR con intervalo distinto del de la tabla",
    "rir_limite_faltante": "RIR sin uno de los límites del T/A",
    "rir_inicio_distinto": "RIR empieza en otra fecha que el T/A",
    "rir_parent_distinto": "RIR en otro Parent Assembly",
    "rir_intervalo_menor": "RIR más seguido que la tabla (más conservador)",
    "asociacion_aproximada": "Fault Name con errores de escritura",
    "asociacion_por_referencia_tecnica": "Asociado por la referencia técnica, no por el nombre",
    "asociacion_por_descripcion": "Asociado por la descripción, no por el nombre",
    "asociacion_por_tabla": "Asociado a una fila de la tabla sin entrada en el catálogo",
    "completado_fuera_de_plazo": "Cerrado después del due date",
    "cierre_deadline_extendido": "Cerrado después del due date original (deadline extendido)",
    "cierre_sin_accion_correctiva": "Cerrado sin acción correctiva",
    "cierre_sin_referencia_tecnica": "Acción correctiva sin referencia del manual",
    "cierre_referencia_incompleta": "Acción correctiva con referencia sin revisión o sin fecha",
    "cierre_sn_distinto": "S/N removido distinto del de la descripción",
    "cierre_sin_adjuntos": "Cerrado sin attachments",
    "sin_adjuntos": "Sin attachments",
    "inicio_ta_incorrecto": "Start del T/A distinto de la apertura (Incorrect DIP Start Values)",
    "sn_inconsistente": "S/N distinto del de las notas de la descripción",
    "referencia_tecnica_no_especifica": "Referencia no específica (barcode de trabajo)",
    "limite_faltante": "Al T/A le falta uno de los límites de la tabla (X o Y)",
    "intervalo_segun_manual": "TA según el daño o el tipo: revisar en el manual",
    "rir_segun_manual": "RIR según el daño o el tipo: revisar en el manual",
    "intervalo_mayor_a_descripcion": "Intervalo mayor que el plazo de la descripción",
    "tipo_intervalo_distinto_a_descripcion": "Tipo de intervalo distinto del de la descripción",
    "limite_declarado_faltante": "Al T/A le falta un límite que pide la descripción",
    "rir_descripcion_faltante": "Falta el RIR que pide la descripción",
    "rir_descripcion_distinto": "RIR distinto del que pide la descripción",
    "parent_assy_distinto_a_pares": "Parent Assembly distinto de otros DIPs con la misma referencia",
    "programado_despues_del_due": "Paquete de trabajo programado después del due date",
    "adjunto_otro_avion": "Attachment con la matrícula de otro avión",
}
COINCIDENCIAS = {  # (auditor con discrepancia, resultado del programa) -> texto
    (True, POSIBLE): "Sí: ambos encuentran discrepancia",
    (True, "OK"): "No: el programa no la detecta",
    (False, "OK"): "Sí: ambos OK",
    (False, POSIBLE): "No: el programa marca posible discrepancia y el auditor no",
    # Sin fuente no es OK ni discrepancia: no se cuenta como acierto ni como error.
    (True, SIN_FUENTE): f"{SIN_FUENTE}: el auditor encontró discrepancia",
    (False, SIN_FUENTE): f"{SIN_FUENTE}: el auditor no encontró discrepancia",
    # Revisar tampoco: el programa no decide, remite al manual.
    (True, REVISAR): "Revisar en el manual: el auditor encontró discrepancia",
    (False, REVISAR): "Revisar en el manual: el auditor no encontró discrepancia",
}
# Las columnas que vienen del Excel de entrada conservan su nombre original; las
# del programa usan los mismos nombres que el reporte de auditoría.
COLUMNAS_DETALLE = [
    "Fault Barcode", "Aircraft", "Fault Name", "Fault Status", "REFERENCIA ASOCIADA",
    COLUMNAS_MANUALES["apertura"], COLUMNAS_MANUALES["auditor_apertura"],
    "Apertura: resultado del programa", "Apertura: motivos",
    COLUMNAS_MANUALES["cierre"], COLUMNAS_MANUALES["auditor_cierre"],
    "Cierre: resultado del programa", "Cierre: motivos",
    COLUMNAS_MANUALES["discrepancia"], "Programa: resultado", "¿Coinciden?", "Leído en ICAM",
    "Fila Excel",
]


@dataclass
class Comparacion:
    lote: Lote
    detalle: pd.DataFrame      # un DIP por fila: auditoría manual y del programa
    hallazgos: pd.DataFrame    # hallazgos del programa en las auditorías hechas por el auditor
    tablas: list[tuple[str, str, pd.DataFrame]] = field(default_factory=list)  # resumen


def _columnas_manuales(original: pd.DataFrame) -> dict[str, str]:
    minusculas = {c.lower().strip(): c for c in original.columns}
    faltan = [n for n in COLUMNAS_MANUALES.values() if n.lower() not in minusculas]
    if faltan:
        raise ErrorCarga("El Excel no trae las columnas de la auditoría manual: " + ", ".join(faltan))
    return {k: minusculas[n.lower()] for k, n in COLUMNAS_MANUALES.items()}


def _resultado(estado: str) -> str:
    """Estado del programa -> resultado mostrado: OK, sin fuente, revisar, posible discrepancia o ""
    (no auditado)."""
    return {"": "", "OK": "OK", SIN_FUENTE: SIN_FUENTE, "REVISAR": REVISAR}.get(estado or "", POSIBLE)


def _peor(resultados: list[str]) -> str:
    resultados = [r for r in resultados if r]
    return max(resultados, key=RESULTADOS.index) if resultados else ""


def comparar(auditor: Auditor, ruta: Path, hoja: str | None = None,
             icam: Path | None = None) -> Comparacion:
    lote = cargar_export(ruta, hoja)
    lote.prueba = bool(auditor.config.get("modo_prueba"))
    if icam:
        completar_con_icam(lote, icam)
    col = _columnas_manuales(lote.original)
    manual = pd.DataFrame({
        k: [parse_fecha(v) if k in ("apertura", "cierre") else parse_texto(v)
            for v in lote.original[c]] for k, c in col.items()}, index=lote.original.index,
        dtype=object)

    # Dos pasadas: a la hora de la auditoría de apertura y a la de cierre.
    apertura = auditor.auditar_lote(lote, list(manual["apertura"]))
    cierre = auditor.auditar_lote(lote, list(manual["cierre"]))
    fa = apertura.filas.set_index("Fila Excel")
    fc = cierre.filas.set_index("Fila Excel")

    # Solo los DIPs en alcance con alguna auditoría manual; de cada pasada, los hallazgos
    # de la auditoría que corresponde.
    en_alcance = set(fa.index[fa["ESTADO IA"] != FUERA_DE_ALCANCE])
    con_apertura = {f for f in manual.index[manual["apertura"].notna()] if f in en_alcance}
    con_cierre = {f for f in manual.index[manual["cierre"].notna()] if f in en_alcance}
    ha, hc = apertura.hallazgos, cierre.hallazgos
    # Sin fuente vale también para el cierre: en los DIPs que el auditor solo revisó al
    # cerrar, es el motivo de ese resultado.
    sin_fuente_cierre = hc[(hc["Severidad"] == SEVERIDAD_SIN_FUENTE)
                           & hc["Fila Excel"].isin(con_cierre - con_apertura)]
    h = pd.concat([ha[(ha["Auditoría"] == "APERTURA") & ha["Fila Excel"].isin(con_apertura)],
                   hc[(hc["Auditoría"] == "CIERRE") & hc["Fila Excel"].isin(con_cierre)],
                   sin_fuente_cierre.assign(**{"Auditoría": "CIERRE"})])
    h = h.sort_values(["Fila Excel", "Auditoría"], ascending=[True, False])
    h.insert(h.columns.get_loc("Código"), "Motivo", h["Código"].map(MOTIVOS).fillna(h["Código"]))
    motivos = {(f, e): "\n".join(f"- {d}" for s, d in zip(g["Severidad"], g["Detalle"])
                                 if s != "OBSERVACION")
               for (f, e), g in h.groupby(["Fila Excel", "Auditoría"])}

    leido = lote.datos.set_index("fila")["adjuntos"].map(lambda x: x is not None)
    filas = []
    for fila_excel in sorted(con_apertura | con_cierre):
        a, m = fa.loc[fila_excel], manual.loc[fila_excel]
        estado_ap = a["ESTADO APERTURA"] if fila_excel in con_apertura else ""
        estado_ci = fc.loc[fila_excel, "ESTADO CIERRE"] if fila_excel in con_cierre else ""
        programa = _peor([_resultado(estado_ap), _resultado(estado_ci)])
        disc_manual = m["discrepancia"].upper() == "Y"
        filas.append([
            a["Fault Barcode"], a["Aircraft"], a["Fault Name"], a["Fault Status"],
            a["REFERENCIA ASOCIADA"] or "(no encontrada)",
            m["apertura"], m["auditor_apertura"], _resultado(estado_ap),
            motivos.get((fila_excel, "APERTURA"), ""),
            m["cierre"], m["auditor_cierre"], _resultado(estado_ci),
            motivos.get((fila_excel, "CIERRE"), ""),
            m["discrepancia"], programa,
            COINCIDENCIAS[(disc_manual, programa)], "Sí" if leido[fila_excel] else "No",
            fila_excel,
        ])
    detalle = pd.DataFrame(filas, columns=COLUMNAS_DETALLE)
    return Comparacion(lote=lote, detalle=detalle, hallazgos=h,
                       tablas=_tablas_resumen(detalle, h, len(con_apertura), len(con_cierre)))


def _pct(a: int, b: int) -> str:
    return f"{a / b * 100:.0f} %" if b else "-"


def _tablas_resumen(d: pd.DataFrame, h: pd.DataFrame, n_apertura: int,
                    n_cierre: int) -> list[tuple[str, str, pd.DataFrame]]:
    """Tablas del resumen: cada una responde una sola pregunta, con una fila por respuesta."""
    total = len(d)
    coincide = d["¿Coinciden?"]
    disc = d[COLUMNAS_MANUALES["discrepancia"]].str.upper() == "Y"
    con_disc = d[disc]
    sin_disc = d[~disc]

    def cuenta(serie: pd.Series) -> int:
        return int(serie.sum())

    def lista(filas: list[tuple], columnas: list[str]) -> pd.DataFrame:
        return pd.DataFrame(filas, columns=columnas)

    tablas = []
    leidos = cuenta(d["Leído en ICAM"] == "Sí")
    sin_tabla = cuenta(d["REFERENCIA ASOCIADA"] == "(no encontrada)")
    tablas.append(("Qué se comparó",
                   "Cada DIP se audita a la fecha y hora en que lo auditó el auditor, en la apertura "
                   "y en el cierre.",
                   lista([("DIPs comparados", total, ""),
                          ("Auditorías de apertura", n_apertura, ""),
                          ("Auditorías de cierre", n_cierre, ""),
                          ("DIPs leídos en ICAM", leidos, _pct(leidos, total)),
                          ("DIPs sin fuente (no están en la tabla de referencias)", sin_tabla,
                           _pct(sin_tabla, total))],
                         ["Dato", "Cantidad", "%"])))

    iguales, distintos = cuenta(coincide.str.startswith("Sí")), cuenta(coincide.str.startswith("No"))
    sin_fuente = cuenta(coincide.str.startswith(SIN_FUENTE))
    revisar = cuenta(coincide.str.startswith(REVISAR))
    tablas.append(("¿El programa llega al mismo resultado que el auditor?",
                   "El programa marca posible discrepancia si alguna de las auditorías del DIP la tiene. "
                   f"{SIN_FUENTE}: el DIP no está en la tabla de referencias y no se compara. "
                   "Revisar en el manual: la tabla da el TA o el RIR según el daño y no se compara.",
                   lista([("Sí, mismo resultado", iguales, _pct(iguales, total)),
                          ("No, resultado distinto", distintos, _pct(distintos, total)),
                          (SIN_FUENTE, sin_fuente, _pct(sin_fuente, total)),
                          (REVISAR, revisar, _pct(revisar, total))],
                         ["Respuesta", "DIPs", "%"])))

    def segun_programa(grupo: pd.DataFrame, lecturas: dict[str, str]) -> pd.DataFrame:
        conteo = grupo["Programa: resultado"].value_counts()
        return lista([(r, lecturas[r], int(conteo.get(r, 0)), _pct(int(conteo.get(r, 0)), len(grupo)))
                      for r in RESULTADOS],
                     ["Resultado del programa", "Qué significa", "DIPs", "%"])

    tablas.append((f"DIPs que el auditor dio sin discrepancia ({len(sin_disc)})",
                   "Lo ideal: que el programa también dé OK.",
                   segun_programa(sin_disc, {
                       "OK": "Coincide con el auditor",
                       SIN_FUENTE: "No está en la tabla de referencias",
                       REVISAR: "La tabla remite al manual",
                       POSIBLE: "El auditor no la marcó"})))
    tablas.append((f"DIPs en que el auditor encontró discrepancia ({len(con_disc)})",
                   "Lo ideal: que el programa también marque discrepancia.",
                   segun_programa(con_disc, {
                       "OK": "No la detecta",
                       SIN_FUENTE: "No está en la tabla de referencias",
                       REVISAR: "La tabla remite al manual",
                       POSIBLE: "Coincide con el auditor"})))

    for nombre, columna in (("apertura", "Apertura: resultado del programa"),
                            ("cierre", "Cierre: resultado del programa")):
        conteo = d[columna].value_counts()
        auditadas = int(sum(conteo.get(r, 0) for r in RESULTADOS))
        tablas.append((f"Auditoría de {nombre}: resultado del programa ({auditadas} auditorías)", "",
                       lista([(r, int(conteo.get(r, 0)), _pct(int(conteo.get(r, 0)), auditadas))
                              for r in RESULTADOS],
                             ["Resultado del programa", "Auditorías", "%"])))

    marcados = set(sin_disc.loc[sin_disc["Programa: resultado"] == POSIBLE, "Fault Barcode"])
    for auditoria in ("APERTURA", "CIERRE"):
        x = h[(h["Auditoría"] == auditoria) & h["Fault Barcode"].isin(marcados)
              & h["Severidad"].isin(["CRITICA", "NO CRITICA", "REVISAR"])]
        filas = []
        for motivo, g in x.groupby("Motivo"):
            dips = list(g["Fault Barcode"].unique())
            filas.append((motivo, len(dips), ", ".join(dips[:3]) + (" …" if len(dips) > 3 else "")))
        tablas.append((f"Por qué el programa marca DIPs que el auditor dio sin discrepancia: "
                       f"{auditoria.lower()}",
                       "Un DIP puede tener más de un motivo.",
                       lista(filas, ["Motivo", "DIPs", "Ejemplos"])
                       .sort_values("DIPs", ascending=False)))

    tablas.append(("Discrepancias encontradas por el auditor",
                   "El Excel no dice qué encontró el auditor; se muestra qué encontró el programa.",
                   con_disc[["Fault Barcode", "Fault Name", "Fault Status", "Programa: resultado",
                             "Apertura: motivos", "Cierre: motivos"]].reset_index(drop=True)))
    return tablas


# ------------------------------------------------------------------ Excel
_AZUL = PatternFill("solid", fgColor="1F3A5F")
_MORADO = PatternFill("solid", fgColor="5B3A8C")  # columnas agregadas por el programa
_ORIGINALES = {"Fault Barcode", "Aircraft", "Fault Name", "Fault Status",
               *COLUMNAS_MANUALES.values()}
_BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
_COLORES = {
    "OK": "C9E7C9", POSIBLE: "FCE3B0", SIN_FUENTE: "D6E4F0", REVISAR: "FFF59D",
    "Observación": "EFEFEF",
    "Sí: ambos OK": "C9E7C9", "Sí: ambos encuentran discrepancia": "C9E7C9",
    "No: el programa no la detecta": "F4B6B6",
    "No: el programa marca posible discrepancia y el auditor no": "FCE3B0",
    f"{SIN_FUENTE}: el auditor encontró discrepancia": "D6E4F0",
    f"{SIN_FUENTE}: el auditor no encontró discrepancia": "D6E4F0",
    "Revisar en el manual: el auditor encontró discrepancia": "FFF59D",
    "Revisar en el manual: el auditor no encontró discrepancia": "FFF59D",
}
_ANCHOS = {"Fault Name": 42, "Apertura: motivos": 70, "Cierre: motivos": 60,
           "Programa: resultado": 32, "Apertura: resultado del programa": 32,
           "Cierre: resultado del programa": 32,
           "REFERENCIA ASOCIADA": 34, "¿Coinciden?": 34, "Motivo": 58, "Ejemplos": 42,
           "Detalle": 90, "Dato": 44, "Respuesta": 44, "Resultado del programa": 26,
           "Qué significa": 48, "Fault Name": 42}


def _valor(v: object) -> object:
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    return v.to_pydatetime() if isinstance(v, pd.Timestamp) else v


def _escribir(ws, fila: int, df: pd.DataFrame, columna: int = 1) -> int:
    """Escribe la tabla con encabezado desde (fila, columna). Devuelve la fila siguiente."""
    for j, nombre in enumerate(df.columns, start=columna):
        c = ws.cell(fila, j, nombre)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), _AZUL, _BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, registro in enumerate(df.itertuples(index=False), start=fila + 1):
        for j, v in enumerate(registro, start=columna):
            c = ws.cell(i, j, _valor(v))
            c.border = _BORDE
            c.alignment = Alignment(wrap_text=isinstance(v, str) and len(v) > 30, vertical="top")
            if isinstance(c.value, datetime):
                c.number_format = "dd-mmm-yyyy hh:mm"
            if isinstance(v, str) and v in _COLORES:
                c.fill = PatternFill("solid", fgColor=_COLORES[v])
    return fila + len(df) + 1


def _anchos(ws, columnas: list[str], columna: int = 1, minimo: int = 14) -> None:
    for j, nombre in enumerate(columnas, start=columna):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = max(ws.column_dimensions[letra].width or 0,
                                                _ANCHOS.get(nombre, minimo))


def exportar(c: Comparacion, carpeta: Path) -> Path:
    carpeta.mkdir(parents=True, exist_ok=True)
    prefijo = "PRUEBA_" if c.lote.prueba else ""
    ruta = carpeta / f"{prefijo}{c.lote.ruta.stem}_COMPARACION_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Programa vs. auditoría manual"
    ws["A1"].font = Font(bold=True, size=16, color="1F3A5F")
    ws["A2"] = f"{c.lote.ruta.name}, generado el {datetime.now():%d-%b-%Y %H:%M}"
    ws["A2"].font = Font(italic=True, color="666666")
    fila = 4
    if c.lote.prueba:
        ws["A3"] = "PRUEBA: datos de prueba, no es una auditoría real."
        ws["A3"].font = Font(bold=True, color="C00000")
        fila = 5
    for titulo, nota, df in c.tablas:
        ws.cell(fila, 1, titulo).font = Font(bold=True, size=12, color="1F3A5F")
        if nota:
            ws.cell(fila + 1, 1, nota).font = Font(italic=True, color="666666")
            fila += 1
        fila = _escribir(ws, fila + 1, df) + 1
        _anchos(ws, list(df.columns))
    ws.cell(fila, 1, "Limitación: ICAM muestra el estado actual del DIP, no el que vio el auditor "
                     "en ese momento.").font = Font(italic=True, color="666666")

    ws = wb.create_sheet("Detalle por DIP")
    _escribir(ws, 1, c.detalle)
    for celda in ws[1]:
        if celda.value not in _ORIGINALES:
            celda.fill = _MORADO
    _anchos(ws, list(c.detalle.columns))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Hallazgos del programa")
    h = c.hallazgos.assign(Severidad=c.hallazgos["Severidad"].map(SEVERIDADES))
    h = h.rename(columns={"Severidad": "Resultado"})
    h = h[["Fault Barcode", "Aircraft", "Fault Name", "Auditoría", "Resultado", "Motivo", "Detalle",
           "Fila Excel"]]
    _escribir(ws, 1, h)
    _anchos(ws, list(h.columns))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)
    return ruta
