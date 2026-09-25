"""Filtra el Excel "DIPs Audit (...)" y deja solo los DIPs que faltan por auditar.

Se queda con las filas que cumplen alguna de estas reglas:

- Fault Status ACTV (Active) y "DIP Open Audit User" vacío: falta la auditoría de apertura.
- Fault Status COMPLETE y "DIP Closed Audit User" vacío: falta la auditoría de cierre.

Se corre por separado de la auditoría. Guarda una copia junto al original, con " FILTERED" al
final del nombre, y conserva las hojas, el formato y los filtros del archivo.

Uso: python -m dips_audit.filtrar "DIPs Audit (25SEP2026).xlsx"
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

COL_ESTADO = "Fault Status"
COL_OPEN_USER = "DIP Open Audit User"
COL_CLOSED_USER = "DIP Closed Audit User"

ACTIVO = {"ACTV", "ACTIVE"}
COMPLETO = {"COMPLETE"}


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _estado(valor) -> str:
    # ICAM a veces trae "ACTV (Active)": vale el código.
    texto = _texto(valor).upper()
    return texto.split()[0] if texto else ""


def pendiente(estado, open_user, closed_user) -> bool:
    """True si al DIP le falta la auditoría que corresponde a su estado."""
    e = _estado(estado)
    if e in ACTIVO:
        return not _texto(open_user)
    if e in COMPLETO:
        return not _texto(closed_user)
    return False


def _filtrar_hoja(ws) -> tuple[int, int] | None:
    """Borra de la hoja las filas que no cumplen las reglas. None si no tiene las columnas."""
    encabezado = {_texto(c.value): c.column for c in ws[1]}
    if not {COL_ESTADO, COL_OPEN_USER, COL_CLOSED_USER} <= encabezado.keys():
        return None
    ce, co, cc = encabezado[COL_ESTADO], encabezado[COL_OPEN_USER], encabezado[COL_CLOSED_USER]

    borrar = [
        fila for fila in range(2, ws.max_row + 1)
        if not pendiente(ws.cell(fila, ce).value, ws.cell(fila, co).value, ws.cell(fila, cc).value)
    ]
    total = ws.max_row - 1
    # De abajo hacia arriba y por tramos seguidos: delete_rows es lento fila por fila.
    i = len(borrar) - 1
    while i >= 0:
        fin = borrar[i]
        while i > 0 and borrar[i - 1] == borrar[i] - 1:
            i -= 1
        ws.delete_rows(borrar[i], fin - borrar[i] + 1)
        i -= 1

    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    return total - len(borrar), total


def ruta_salida(entrada: Path) -> Path:
    return entrada.with_name(f"{entrada.stem} FILTERED{entrada.suffix}")


def filtrar(entrada: Path) -> tuple[Path, dict[str, tuple[int, int]]]:
    wb = load_workbook(entrada)
    resumen = {}
    for ws in wb.worksheets:
        r = _filtrar_hoja(ws)
        if r is not None:
            resumen[ws.title] = r
    if not resumen:
        raise ValueError(
            f"Ninguna hoja tiene las columnas {COL_ESTADO!r}, {COL_OPEN_USER!r} y {COL_CLOSED_USER!r}."
        )
    salida = ruta_salida(entrada)
    wb.save(salida)
    return salida, resumen


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m dips_audit.filtrar",
        description="Deja solo los DIPs ACTV sin auditoría de apertura y COMPLETE sin auditoría de cierre.",
    )
    parser.add_argument("excel", type=Path, help='Excel "DIPs Audit (...)"')
    args = parser.parse_args(argv)

    if not args.excel.is_file():
        print(f"No existe el archivo: {args.excel}", file=sys.stderr)
        return 1
    try:
        salida, resumen = filtrar(args.excel)
    except (ValueError, OSError) as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1
    for hoja, (quedan, total) in resumen.items():
        print(f"{hoja}: {quedan} de {total} filas")
    print(f"Guardado: {salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
